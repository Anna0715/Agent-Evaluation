#!/usr/bin/env python3
"""补充关键字模糊匹配用例：打错字 / 没打全仍应命中对应周报事实。

复用已有 VEC-RICH / VEC-KB / ROLE-PERIOD 语料，不新发报。
用例：WA-134 ~ WA-139。
"""

from __future__ import annotations

from _paths import CATALOG_PATH, FIXTURES_DIR, ROLE_PERIOD_PEOPLE_PATH, XLSX_PATH
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

CASE_HEADER_ROW = 2
RECORD_HEADER_ROW = 2
RECEIVER = "智本_anrou"
COMPANY = "智本科技"
EXPECT_TOOLS = (
    '{"required":["memory_search"],"forbidden":["memory_settle"],"bound_entity_only":true}'
)

PRE_VEC = (
    f"追问账号={RECEIVER}；追问公司={COMPANY}。"
    "已有霜灯索引长周报（VEC-RICH，召回@10=0.917，QPS 1480）、"
    "潮汐对账长周报（银杏轧差错账率 0.37%）、"
    "林秋禾回访周报（RD-8831，满意度 4.6）、"
    "ZX-Cloud-Lite 发版周报（PAY-TIMEOUT-17，灰度 12%）、"
    "ROLE-PERIOD 青检台/灰灯回归周报。"
    "不附带周报，由 Agent 做关键词模糊检索。"
)

CASES: list[dict[str, str]] = [
    {
        "用例ID": "WA-134",
        "一级模块": "自然语言选周报",
        "场景": "错别字检索",
        "优先级": "P1",
        "优先级说明": "项目名打错一字仍应命中霜灯索引长周报",
        "测试输入/用户话术": "帮我找霜灯索隐相关的周报，召回@10是多少？",
        "前置条件": PRE_VEC + "正确写法是「霜灯索引」。",
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送错字「霜灯索隐」问法；\n"
            "4. 应落到霜灯索引周报并给出召回@10=0.917。"
        ),
        "必须满足": "霜灯索引；0.917",
        "禁止出现": "银杏轧差；潮汐对账；0.37%；找不到；未命中",
        "预期结果": "错别字「霜灯索隐」仍命中霜灯索引周报，写出召回@10=0.917。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "向量匹配/错别字",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-135",
        "一级模块": "自然语言选周报",
        "场景": "错别字检索",
        "优先级": "P1",
        "优先级说明": "人名打错一字仍应命中客户回访周报",
        "测试输入/用户话术": "林秋和那单回访得怎么样了？",
        "前置条件": PRE_VEC + "正确写法是「林秋禾」。",
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 用错字「林秋和」提问；\n"
            "4. 应答出 RD-8831 与满意度 4.6。"
        ),
        "必须满足": "林秋禾；RD-8831；4.6",
        "禁止出现": "HEALTH-9921；0.37%；霜灯索引；找不到",
        "预期结果": "错别字「林秋和」仍命中林秋禾回访，订单 RD-8831，满意度 4.6。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "向量知识库/错别字",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-136",
        "一级模块": "自然语言选周报",
        "场景": "关键词缺字检索",
        "优先级": "P1",
        "优先级说明": "只打项目名前两字，仍应补全命中霜灯索引",
        "测试输入/用户话术": "霜灯这周召回效果怎么样？召回@10多少？",
        "前置条件": PRE_VEC + "用户未打全「霜灯索引」。",
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送缺字问法「霜灯」；\n"
            "4. 应补全到霜灯索引并给出 0.917。"
        ),
        "必须满足": "霜灯索引；0.917",
        "禁止出现": "银杏轧差；0.37%；找不到相关周报",
        "预期结果": "关键词未打全「霜灯」仍命中霜灯索引周报，写出召回@10=0.917。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "向量匹配/缺字",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-137",
        "一级模块": "自然语言选周报",
        "场景": "关键词缺字检索",
        "优先级": "P1",
        "优先级说明": "项目代号只打前缀，仍应命中 FROST-LANTERN-IDX",
        "测试输入/用户话术": "搜一下 FROST-LANTERN 相关周报，QPS 提到多少了？",
        "前置条件": PRE_VEC + "完整代号是 FROST-LANTERN-IDX。",
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送不完整代号 FROST-LANTERN；\n"
            "4. 应命中霜灯索引周报并给出 QPS=1480。"
        ),
        "必须满足": "FROST-LANTERN-IDX；1480",
        "禁止出现": "0.37%；银杏轧差；找不到",
        "预期结果": "缺字代号 FROST-LANTERN 仍命中 FROST-LANTERN-IDX，QPS=1480。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "向量匹配/缺字代号",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-138",
        "一级模块": "自然语言选周报",
        "场景": "错别字检索",
        "优先级": "P1",
        "优先级说明": "对账项目名错字仍应命中潮汐对账，不串霜灯",
        "测试输入/用户话术": "银杏轧蹉这周错账率是多少？",
        "前置条件": PRE_VEC + "正确写法是「银杏轧差」。",
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 用错字「银杏轧蹉」提问；\n"
            "4. 应给出错账率 0.37%，不得串写霜灯索引。"
        ),
        "必须满足": "0.37%；银杏轧差或潮汐对账",
        "禁止出现": "霜灯索引；0.917；1480；青岚实验舱",
        "预期结果": "错别字「银杏轧蹉」仍命中潮汐对账周报，错账率=0.37%，不串霜灯。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "向量匹配/错别字",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-139",
        "一级模块": "自然语言选周报",
        "场景": "关键词缺字检索",
        "优先级": "P1",
        "优先级说明": "错误码前缀未打全，仍应命中发版周报灰度",
        "测试输入/用户话术": "PAY-TIMEOUT 是什么问题？现在灰度到多少了？",
        "前置条件": PRE_VEC + "完整错误码是 PAY-TIMEOUT-17。",
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送缺字错误码 PAY-TIMEOUT；\n"
            "4. 应补全到 PAY-TIMEOUT-17，并给出 ZX-Cloud-Lite 灰度 12%。"
        ),
        "必须满足": "PAY-TIMEOUT-17；ZX-Cloud-Lite；12%",
        "禁止出现": "4.6；RD-8831；HEALTH-9921；找不到",
        "预期结果": "缺字「PAY-TIMEOUT」仍命中 PAY-TIMEOUT-17 与灰度 12%。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "向量知识库/缺字",
        "expect_tools": EXPECT_TOOLS,
    },
]


def _sheet_headers(ws: Worksheet, header_row: int) -> list[str]:
    return [
        str(cell.value or "").strip()
        for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row))
    ]


def _copy_row_style(ws: Worksheet, src_row: int, dst_row: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)
            dst.alignment = copy(src.alignment)


def write_xlsx(path: Path = XLSX_PATH) -> None:
    wb = load_workbook(path)
    case_ws = wb["用例库"]
    record_ws = wb["执行记录"]
    case_headers = _sheet_headers(case_ws, CASE_HEADER_ROW)
    record_headers = _sheet_headers(record_ws, RECORD_HEADER_ROW)
    case_idx = {name: i + 1 for i, name in enumerate(case_headers) if name}
    rec_case_col = record_headers.index("用例ID") + 1 if "用例ID" in record_headers else None
    existing_rec = set()
    if rec_case_col:
        for row in range(RECORD_HEADER_ROW + 1, record_ws.max_row + 1):
            cid = str(record_ws.cell(row, rec_case_col).value or "").strip()
            if cid:
                existing_rec.add(cid)

    last_case_row = CASE_HEADER_ROW
    case_rows: dict[str, int] = {}
    for row in range(CASE_HEADER_ROW + 1, case_ws.max_row + 1):
        cid = str(case_ws.cell(row, case_idx["用例ID"]).value or "").strip()
        if cid:
            last_case_row = row
            case_rows[cid] = row

    for spec in CASES:
        cid = spec["用例ID"]
        if cid in case_rows:
            target = case_rows[cid]
            print(f"[UPDATE] {cid} {spec['场景']}")
        else:
            last_case_row += 1
            _copy_row_style(case_ws, last_case_row - 1, last_case_row, len(case_headers))
            target = last_case_row
            case_rows[cid] = target
            print(f"[CASE] {cid} {spec['场景']}")
        for key, value in spec.items():
            if key in case_idx:
                case_ws.cell(target, case_idx[key], value=value)
        if rec_case_col and cid not in existing_rec:
            last_rec = RECORD_HEADER_ROW
            for row in range(RECORD_HEADER_ROW + 1, record_ws.max_row + 1):
                if str(record_ws.cell(row, rec_case_col).value or "").strip():
                    last_rec = row
            last_rec += 1
            _copy_row_style(record_ws, last_rec - 1, last_rec, len(record_headers))
            record_ws.cell(last_rec, rec_case_col, value=cid)
            existing_rec.add(cid)

    wb.save(path)
    print(f"[XLSX] WA-134..139 -> {path.name}")


def write_catalog(path: Path = CATALOG_PATH) -> None:
    data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    notes = data.setdefault("case_notes", {})
    notes.update(
        {
            "WA-134": "错字「霜灯索隐」仍命中霜灯索引召回@10=0.917",
            "WA-135": "错字「林秋和」仍命中林秋禾回访 RD-8831 / 4.6",
            "WA-136": "缺字「霜灯」仍命中霜灯索引召回@10=0.917",
            "WA-137": "缺字代号 FROST-LANTERN 仍命中 IDX，QPS=1480",
            "WA-138": "错字「银杏轧蹉」仍命中潮汐对账错账率 0.37%",
            "WA-139": "缺字 PAY-TIMEOUT 仍命中 PAY-TIMEOUT-17 灰度 12%",
        }
    )
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("[CATALOG] wrote WA-134..139")


def main() -> int:
    write_xlsx()
    write_catalog()
    ids = ",".join(item["用例ID"] for item in CASES)
    print(f"[NEXT] python run_report_agent_cases.py --case-ids {ids}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
