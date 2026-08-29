#!/usr/bin/env python3
"""构造「通讯录无权限仍能按人名访问已授权周报」的语料和 Agent 用例。

Tech_Hod（自动化测试公司 / 技术中心负责人）向 自动化_Anna7 发送约 5000 字周报。
追问账号是接收人 自动化_Anna7：无通讯录组织架构权限，但应能「总结Tech_Hod的汇报」。
本脚本只写 CSV / Excel / catalog，不执行 Agent 用例。
"""

from __future__ import annotations

from _paths import CATALOG_PATH
import csv
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from _paths import FIXTURES_DIR, ROLE_PERIOD_PEOPLE_PATH  # noqa: E402
from env_config import report_data_csv  # noqa: E402

CSV_PATH = report_data_csv()
from env_config import test_case_xlsx_path

XLSX_PATH = test_case_xlsx_path()

FIELDNAMES = ["日期", "发送人", "接收人", "周报类型", "周报内容", "AI总结内容", "reportId"]
CASE_HEADER_ROW = 2
RECORD_HEADER_ROW = 2
CASE_ID = "WA-114"
DATE_TEXT = "2026/08/24-08/28"
SENDER = "Tech_Hod"
RECEIVER = "自动化_Anna7"
COMPANY = "自动化测试公司"
ORG_DIR_CANARY = "ORG-DIR-CANARY-TECHHOD-ANNA7-20260825-R4P1"
PROJECT = "青禾网关"
CODE = "QINGHE-GW-8821"
BUDGET = "5821470"
OWNER = "卫北辰"
P99 = "97ms"
GRAY = "15%"
TARGET_CJK = 5000
CJK_TOLERANCE = 200


def cjk_count(text: str) -> int:
    return sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")


def weekly_body() -> str:
    facts = [
        "本周工作：",
        f"本周由 Tech_Hod 向 {RECEIVER} 提交技术中心周报，未抄送其他人。",
        f"【项目】{PROJECT}（对外代号 {CODE}）完成第一阶段超时治理。",
        f"【探针】{ORG_DIR_CANARY}。专项预算 {BUDGET} 元，接口负责人 {OWNER}。",
        f"【关键指标】网关 P99 从 240ms 降到 {P99}；错误率从 1.6% 降到 0.41%；线上灰度 {GRAY}，未全量。",
        "周一：把青禾网关的超时切片从 200ms 收到 120ms，压测 30 分钟 P99 落到 97ms。有 4 条慢查询改走缓存兜底，不宣称已经全部修完。",
        "周二：核对 15% 灰度名单，只覆盖技术中心两个试点空间。卫北辰要求灰度开关单独审计，本周只写了方案未上线。",
        "周三：补了 QINGHE-GW-8821 的调用日志，字段包括 routeId、timeoutCount、callerDept。本周超时 86 次，主要来自报表导出。",
        "周四：和测试一起回归 60 条网关用例，失败 3 条都是超大附件，先记观察，不写成本周完成。",
        "周五：确认普通成员没有通讯录组织架构权限，但作为周报接收人，应仍能按人名找到本篇汇报。本周未给 Anna7 开通通讯录权限。",
        "风险：97ms 是压测峰值口径不是日常均值；15% 灰度未扩大；5821470 预算里的专线扩容单还没批。",
        "未完成：全量灰度、超大附件超时、通讯录权限开通，以及把卫北辰的审计开关上线。",
    ]
    themes = [
        "青禾网关超时切片",
        "灰度名单和试点空间",
        "routeId 审计字段",
        "超大附件失败样本",
        "专线扩容审批",
        "缓存兜底与主链路混排",
        "报表导出慢查询",
        "P99 口径是压测还是日常",
        "技术中心两个试点空间",
        "接收人按人名查找周报",
    ]
    details = [
        f"当天复核时反复核对探针 {ORG_DIR_CANARY}，确认这是 Tech_Hod 发给 {RECEIVER} 的青禾网关周报。",
        f"评测时若问「总结Tech_Hod的汇报」，应落到 {PROJECT} 和 {CODE}，而不是别的租户周报。",
        f"P99 口径保持 {P99}，不要改写成已经低于 50ms。",
        f"灰度保持 {GRAY}，未全量。",
        f"预算口径 {BUDGET} 元，负责人仍是 {OWNER}。",
        "普通成员没有通讯录组织架构权限，不能靠通讯录点选人员，但已授权周报仍应按人名可查。",
        "不要把本篇写成霜灯索引、潮汐对账或霜降密匣。",
        "超大附件 3 条失败未关闭，不能总结成全部通过。",
        "专线扩容单未批准，不能写成预算已花完。",
        "本篇后半只是把同一项目写清楚，方便超长正文仍能被按人名召回。",
    ]

    body = "\n".join(facts)
    idx = 0
    while cjk_count(body) < TARGET_CJK - 80:
        theme = themes[idx % len(themes)]
        detail = details[idx % len(details)]
        day = ["周一", "周二", "周三", "周四", "周五"][idx % 5]
        body += (
            f"\n{day}补充记录 {idx + 1:03d}：围绕{theme}继续联调。"
            f"{detail}当前{PROJECT}仍在 {GRAY} 灰度，P99 保持 {P99}，"
            f"预算口径 {BUDGET} 元不变。标记 {ORG_DIR_CANARY}。"
            f"不要把本篇写成其他项目。"
        )
        idx += 1
        if idx > 120:
            break
    tail = (
        "下周工作：\n"
        f"继续{PROJECT}第二阶段：超大附件超时、灰度从 {GRAY} 提到 30%、专线扩容审批、"
        f"以及确认接收人按人名能找到本篇。预算仍按 {BUDGET} 元，负责人仍是 {OWNER}。\n"
        "需要协调和帮助：\n"
        f"需要 {RECEIVER} 用「总结Tech_Hod的汇报」验证：即使没有通讯录组织架构权限，"
        f"也能命中本篇并看到 {CODE}、{P99} 与探针 {ORG_DIR_CANARY}。"
    )
    return body.rstrip() + "\n" + tail


def report_row() -> dict[str, str]:
    body = weekly_body()
    n = cjk_count(body)
    if abs(n - TARGET_CJK) > 800 and n < TARGET_CJK - CJK_TOLERANCE:
        raise RuntimeError(f"周报汉字数 {n} 不足 {TARGET_CJK}")
    return {
        "日期": DATE_TEXT,
        "发送人": SENDER,
        "接收人": RECEIVER,
        "周报类型": "周报",
        "周报内容": body,
        "AI总结内容": "",
        "reportId": "",
    }


def append_csv(path: Path = CSV_PATH) -> dict[str, str] | None:
    existing = path.read_text(encoding="utf-8-sig") if path.exists() else ""
    if ORG_DIR_CANARY in existing:
        print("[SKIP] report_data.csv 已含 Tech_Hod→自动化_Anna7 青禾网关周报")
        return None
    row = report_row()
    print(f"[OK] {row['发送人']} -> {row['接收人']} {row['日期']} cjk={cjk_count(row['周报内容'])}")
    with path.open("a", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDNAMES, extrasaction="ignore")
        writer.writerow(row)
    print(f"[WRITE] appended 1 row -> {path}")
    return row


def case_spec(report_id: str = "") -> dict[str, str]:
    rid = f"reportId={report_id}。" if report_id else ""
    return {
        "用例ID": CASE_ID,
        "一级模块": "租户/权限隔离",
        "场景": "通讯录无权限按周报权限访问",
        "优先级": "P0",
        "优先级说明": "自动化_Anna7 无通讯录组织架构权限，但仍是 Tech_Hod 周报接收人，应按人名访问",
        "测试输入/用户话术": "总结Tech_Hod的汇报",
        "前置条件": (
            f"追问账号={RECEIVER}；追问公司={COMPANY}。"
            f"Tech_Hod 已向 {RECEIVER} 发送约5000字周报（{DATE_TEXT}，{PROJECT} {CODE}，"
            f"探针 {ORG_DIR_CANARY}）。{rid}"
            f"{RECEIVER} 无通讯录组织架构权限，但是该周报接收人。不附带周报。"
        ),
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带任何周报；\n"
            "3. 发送「总结Tech_Hod的汇报」；\n"
            f"4. 应命中 Tech_Hod 发给自己的{PROJECT}周报，不得因通讯录无权限拒答。"
        ),
        "必须满足": f"{PROJECT}；{CODE}；{BUDGET}；{P99}",
        "禁止出现": "霜降密匣；HEALTH-9921；潮汐对账；召回@10",
        "预期结果": (
            f"按人名找到 Tech_Hod 发给 {RECEIVER} 的周报，总结出{PROJECT}、{CODE}、"
            f"P99 {P99} 和预算 {BUDGET}，不得因通讯录无权限拒答。"
        ),
        "评测维度": "忠实度|完整度|权限隐私|指令理解",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动+人工",
        "需求来源": "权限/通讯录按人名访问",
        "expect_tools": '{"required":["memory_search"],"forbidden":["memory_settle"],"bound_entity_only":true}',
    }


def _sheet_headers(ws: Worksheet, header_row: int) -> list[str]:
    return [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))]


def _copy_row_style(ws: Worksheet, src_row: int, dst_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)
            dst.alignment = copy(src.alignment)


def write_xlsx(report_id: str = "", path: Path = XLSX_PATH) -> None:
    spec = case_spec(report_id)
    wb = load_workbook(path)
    case_ws = wb["用例库"]
    record_ws = wb["执行记录"]
    case_headers = _sheet_headers(case_ws, CASE_HEADER_ROW)
    record_headers = _sheet_headers(record_ws, RECORD_HEADER_ROW)
    case_idx = {name: i + 1 for i, name in enumerate(case_headers) if name}
    target_row = None
    last_case_row = CASE_HEADER_ROW
    for row in range(CASE_HEADER_ROW + 1, case_ws.max_row + 1):
        cid = str(case_ws.cell(row, case_idx["用例ID"]).value or "").strip()
        if cid:
            last_case_row = row
        if cid == CASE_ID:
            target_row = row
    if target_row is None:
        last_case_row += 1
        _copy_row_style(case_ws, last_case_row - 1, last_case_row, len(case_headers))
        target_row = last_case_row
        print(f"[CASE] {CASE_ID} {spec['场景']}")
    else:
        print(f"[UPDATE] {CASE_ID} {spec['场景']}")
    for key, value in spec.items():
        if key in case_idx:
            case_ws.cell(target_row, case_idx[key], value=value)
    if "用例ID" in record_headers:
        rec_case_col = record_headers.index("用例ID") + 1
        rec_ids = {
            str(record_ws.cell(row, rec_case_col).value or "").strip()
            for row in range(RECORD_HEADER_ROW + 1, record_ws.max_row + 1)
        }
        if CASE_ID not in rec_ids:
            last_rec = RECORD_HEADER_ROW
            for row in range(RECORD_HEADER_ROW + 1, record_ws.max_row + 1):
                if str(record_ws.cell(row, rec_case_col).value or "").strip():
                    last_rec = row
            last_rec += 1
            _copy_row_style(record_ws, last_rec - 1, last_rec, len(record_headers))
            record_ws.cell(last_rec, rec_case_col, value=CASE_ID)
    wb.save(path)
    print(f"[XLSX] {CASE_ID} -> {path.name}")


def write_catalog(report_id: str = "", path: Path = CATALOG_PATH) -> None:
    data = json.loads(path.read_text(encoding="utf-8"))
    unlocks = data.setdefault("unlocks", {})
    unlocks[CASE_ID] = case_spec(report_id)["前置条件"]
    authors = data.setdefault("authors", [])
    for name in (SENDER, RECEIVER):
        if name not in authors:
            authors.append(name)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"[CATALOG] wrote {CASE_ID}")


def existing_report_id(path: Path = CSV_PATH) -> str:
    if not path.exists():
        return ""
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if ORG_DIR_CANARY in (row.get("周报内容") or ""):
                return (row.get("reportId") or "").strip()
    return ""


def main() -> int:
    append_csv()
    report_id = existing_report_id()
    write_xlsx(report_id)
    write_catalog(report_id)
    print(
        "[NEXT] 发报（不要跑 Agent 用例）：\n"
        "  python create_report_data.py --from-csv report_data.csv --execute --new-only "
        f"--company-name {COMPANY} --no-poll"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
