import os
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

_ONLINE_DIR = Path(__file__).resolve().parents[1] / "online"
if str(_ONLINE_DIR) not in sys.path:
    sys.path.insert(0, str(_ONLINE_DIR))

from env_config import DATA_XLSX_PATH, data_sheet_name, data_xlsx_path, get_profile  # noqa: E402


@dataclass
class Account:
    user_name: str
    user_email: str
    user_role: str
    user: str = ""
    company_name: str = ""
    tenant_id: str = ""
    company_email: str = ""
    department_name: str = ""
    department_id: str = ""
    user_id: str = ""
    password: str = ""

    @property
    def login_account(self) -> str:
        return self.user_email


_ACCOUNTS_CACHE: Dict[tuple[str, str], List[Account]] = {}

DATA_CSV_PATH = DATA_XLSX_PATH  # legacy alias
TEST_CASE_FILE_PATH = os.path.join(os.path.dirname(__file__), "test_case.xlsx")
XLSX_TEST_CASE_SHEET_ALIASES: Dict[str, str] = {}

_ACCOUNT_FIELD_MAP: Dict[str, str] = {
    "UserName": "user_name",
    "UserEmail": "user_email",
    "用户角色": "user_role",
    "用户": "user",
    "CompanyName": "company_name",
    "TenantId": "tenant_id",
    "tenant_id": "tenant_id",
    "CompanyEmail": "company_email",
    "DepartmentName": "department_name",
    "DepartmentId": "department_id",
    "org_unit_id": "department_id",
    "org_unit_id ": "department_id",
    "UseriD": "user_id",
    "User_Id": "user_id",
    "UserID": "user_id",
    "Password": "password",
    "password": "password",
}


def _row_to_account(row: dict[str, Any]) -> Account | None:
    user_name = str(row.get("UserName") or row.get("user_name") or "").strip()
    user_email = str(row.get("UserEmail") or row.get("user_email") or "").strip()
    if not user_email:
        return None
    kwargs: dict[str, str] = {}
    for src, attr in _ACCOUNT_FIELD_MAP.items():
        val = row.get(src)
        if val is not None and str(val).strip():
            kwargs[attr] = str(val).strip()
    kwargs.setdefault("user_email", user_email)
    kwargs.setdefault("user_name", user_name)
    return Account(**{k: kwargs.get(k, "") for k in Account.__dataclass_fields__})


def _missing_account_hint(
    xlsx_path: str,
    sheet_name: str,
    user_name: str,
    company_name: str,
) -> str:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        return ""
    wb = load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return ""
    ws = wb[sheet_name]
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1):
        row = {headers[i]: ws.cell(r, i + 1).value for i in range(len(headers))}
        name = str(row.get("UserName") or "").strip()
        company = str(row.get("CompanyName") or "").strip()
        if user_name and name != user_name:
            continue
        if company_name and company != company_name:
            continue
        if not name:
            continue
        missing = []
        if not str(row.get("UserEmail") or "").strip():
            missing.append("UserEmail")
        profile = get_profile()
        if profile.otp_mode == "email" and not str(row.get("Password") or "").strip():
            missing.append("Password")
        if missing:
            return f"（已找到行 {name!r}，但缺少 {','.join(missing)}）"
    return ""


def _read_accounts_from_xlsx(xlsx_path: str, sheet_name: str) -> List[Account]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise ImportError("read account xlsx requires openpyxl, please `pip install openpyxl`") from exc

    wb = load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"sheet {sheet_name!r} not found in {xlsx_path}, available={wb.sheetnames}")
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(v or "").strip() for v in next(rows_iter)]
    except StopIteration:
        return []

    accounts: List[Account] = []
    for values in rows_iter:
        row_dict = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
        account = _row_to_account(row_dict)
        if account:
            accounts.append(account)
    return accounts


def get_accounts(
    xlsx_path: str | None = None,
    *,
    sheet_name: str | None = None,
    env: str | None = None,
) -> List[Account]:
    """读取 data.xlsx 指定 sheet 的全量账号（带缓存）。"""
    xlsx_path = os.path.abspath(xlsx_path or str(data_xlsx_path()))
    sheet = (sheet_name or data_sheet_name(env)).strip()
    cache_key = (xlsx_path, sheet)
    if cache_key in _ACCOUNTS_CACHE:
        return _ACCOUNTS_CACHE[cache_key]
    accounts = _read_accounts_from_xlsx(xlsx_path, sheet)
    _ACCOUNTS_CACHE[cache_key] = accounts
    return accounts


def get_account_from_data_csv(
    csv_path: str,
    *,
    user_role: Optional[str] = None,
    user_name: Optional[str] = None,
    company_name: Optional[str] = None,
    sheet_name: str | None = None,
    env: str | None = None,
) -> Account:
    """从 data.xlsx 精确选择一条账号（兼容旧函数名 csv_path 实为 xlsx 路径）。"""
    accounts = get_accounts(csv_path, sheet_name=sheet_name, env=env)
    user_role = (user_role or "").strip()
    user_name = (user_name or "").strip()
    company_name = (company_name or "").strip()
    if not user_role and not user_name and not company_name:
        raise ValueError("please provide `user_role` or `user_name` or `company_name`")

    matches = accounts
    if user_role:
        matches = [a for a in matches if a.user_role == user_role]
    if user_name:
        matches = [a for a in matches if a.user_name == user_name]
    if company_name:
        matches = [a for a in matches if a.company_name == company_name]

    source = f"data.xlsx/{sheet_name or data_sheet_name(env)}"
    if len(matches) == 1:
        return matches[0]
    if len(matches) == 0:
        hint = _missing_account_hint(
            csv_path,
            sheet_name or data_sheet_name(env),
            user_name,
            company_name,
        )
        raise ValueError(
            f"no account matched user_role={user_role}, user_name={user_name}, "
            f"company_name={company_name} in {source}{hint}"
        )
    raise ValueError(
        f"multiple accounts matched user_role={user_role}, user_name={user_name}, "
        f"company_name={company_name}; expected unique"
    )


_FIELD_TO_ATTR = dict(_ACCOUNT_FIELD_MAP)


def data_csv(
    user_role: str,
    field_name: str,
    csv_path: str = DATA_XLSX_PATH,
    *,
    sheet_name: str | None = None,
    env: str | None = None,
) -> str:
    """返回 data.xlsx 某字段值（兼容旧函数名）。"""
    user_role = (user_role or "").strip()
    field_name = (field_name or "").strip()
    if field_name not in _FIELD_TO_ATTR:
        raise ValueError(f"unsupported field_name={field_name}; supported={sorted(_FIELD_TO_ATTR.keys())}")

    accounts = get_accounts(csv_path, sheet_name=sheet_name, env=env)
    matches = [a for a in accounts if a.user_role == user_role]
    if len(matches) == 0:
        raise ValueError(f"no account matched user_role={user_role} in data.xlsx")
    if len(matches) != 1:
        candidates = [(a.user_role, a.user_email) for a in matches]
        raise ValueError(f"user_role={user_role} matched {len(matches)} rows: {candidates}")

    attr = _FIELD_TO_ATTR[field_name]
    return str(getattr(matches[0], attr))


_TEST_CASE_HEADER_MAP: Dict[str, str] = {
    "场景": "scenario",
    "page": "page",
    "前置条件": "precondition",
    "用例名称": "case_name",
    "user_name": "user_name",
    "company_name": "company_name",
    "task_name": "task_name",
    "task_id": "task_id",
    "预期结果": "expected",
    "space_type": "space_type",
    "owner_mode": "owner_mode",
    "owner_id": "owner_id",
    "member_users": "member_users",
    "member_user_ids": "member_user_ids",
    "expect_error": "expect_error",
    "allow_existing": "allow_existing",
}
TEST_CASE_RESULT_COL = "执行结果"
TEST_CASE_REASON_COL = "原因"


def _normalize_case_row(raw_row: Dict[str, Any]) -> Dict[str, str]:
    normalized: Dict[str, str] = {}
    for src_key, target_key in _TEST_CASE_HEADER_MAP.items():
        normalized[target_key] = str(raw_row.get(src_key, "") or "").strip()
    return normalized


def _read_test_cases_from_xlsx(xlsx_path: str, sheet_name: str) -> List[Dict[str, str]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise ImportError("read xlsx test cases requires openpyxl, please `pip install openpyxl`") from exc

    normalized_sheet = (sheet_name or "").strip()
    if not normalized_sheet:
        raise ValueError("sheet_name is required")

    wb = load_workbook(xlsx_path, data_only=True)
    resolved_sheet = XLSX_TEST_CASE_SHEET_ALIASES.get(normalized_sheet, normalized_sheet)
    if resolved_sheet not in wb.sheetnames:
        raise ValueError(f"sheet '{sheet_name}' not found, available={wb.sheetnames}")
    ws = wb[resolved_sheet]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(v or "").strip() for v in next(rows_iter)]
    except StopIteration:
        return []

    rows: List[Dict[str, str]] = []
    for values in rows_iter:
        row_dict = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
        case_name = str(row_dict.get("用例名称", "") or "").strip()
        if not case_name:
            continue
        rows.append(_normalize_case_row(row_dict))
    return rows


def read_test_cases_by_sheet(sheet_name: str, file_path: str = TEST_CASE_FILE_PATH) -> List[Dict[str, str]]:
    abs_path = os.path.abspath(file_path)
    if not os.path.exists(abs_path):
        raise FileNotFoundError(f"test case file not found: {abs_path}")

    ext = os.path.splitext(abs_path)[1].lower()
    if ext in {".xlsx", ".xlsm"}:
        return _read_test_cases_from_xlsx(abs_path, sheet_name)
    raise ValueError(f"unsupported test case file extension: {ext}, only .xlsx/.xlsm are supported")


def _write_test_case_result_to_xlsx(
    xlsx_path: str,
    sheet_name: str,
    case_name: str,
    execution_result: str,
    reason: str,
) -> None:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise ImportError("write xlsx test case requires openpyxl, please `pip install openpyxl`") from exc

    normalized_sheet = (sheet_name or "").strip()
    resolved_sheet = XLSX_TEST_CASE_SHEET_ALIASES.get(normalized_sheet, normalized_sheet)
    if not resolved_sheet:
        raise ValueError("sheet_name is required")

    normalized_case_name = (case_name or "").strip()
    if not normalized_case_name:
        raise ValueError("case_name is required")

    wb = load_workbook(xlsx_path)
    if resolved_sheet not in wb.sheetnames:
        raise ValueError(f"sheet '{sheet_name}' not found, available={wb.sheetnames}")
    ws = wb[resolved_sheet]

    header_values = [str(v or "").strip() for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    if not header_values:
        raise ValueError(f"sheet '{resolved_sheet}' is empty, missing header row")

    if "用例名称" not in header_values:
        raise ValueError(f"sheet '{resolved_sheet}' missing required column: 用例名称")

    if TEST_CASE_RESULT_COL not in header_values:
        header_values.append(TEST_CASE_RESULT_COL)
        ws.cell(row=1, column=len(header_values), value=TEST_CASE_RESULT_COL)
    if TEST_CASE_REASON_COL not in header_values:
        header_values.append(TEST_CASE_REASON_COL)
        ws.cell(row=1, column=len(header_values), value=TEST_CASE_REASON_COL)

    case_col_idx = header_values.index("用例名称") + 1
    result_col_idx = header_values.index(TEST_CASE_RESULT_COL) + 1
    reason_col_idx = header_values.index(TEST_CASE_REASON_COL) + 1

    matched_rows: List[int] = []
    for row_idx in range(2, ws.max_row + 1):
        cell_value = str(ws.cell(row=row_idx, column=case_col_idx).value or "").strip()
        if cell_value == normalized_case_name:
            matched_rows.append(row_idx)

    if len(matched_rows) == 0:
        raise ValueError(f"case not found in test_case: {normalized_case_name}")
    if len(matched_rows) > 1:
        raise ValueError(f"case_name duplicated in test_case: {normalized_case_name}, count={len(matched_rows)}")

    target_row = matched_rows[0]
    ws.cell(row=target_row, column=result_col_idx, value=(execution_result or "").strip().lower())
    ws.cell(row=target_row, column=reason_col_idx, value=(reason or "").strip())
    wb.save(xlsx_path)


def write_test_case_execution_result(
    *,
    sheet_name: str,
    case_name: str,
    execution_result: str,
    reason: str = "",
    file_path: str = TEST_CASE_FILE_PATH,
) -> None:
    abs_path = os.path.abspath(file_path)
    if not os.path.exists(abs_path):
        raise FileNotFoundError(f"test case file not found: {abs_path}")

    ext = os.path.splitext(abs_path)[1].lower()
    if ext in {".xlsx", ".xlsm"}:
        _write_test_case_result_to_xlsx(
            abs_path,
            sheet_name=sheet_name,
            case_name=case_name,
            execution_result=execution_result,
            reason=reason,
        )
        return
    raise ValueError(f"unsupported test case file extension: {ext}, only .xlsx/.xlsm are supported")
