#!/usr/bin/env python3
"""补充正确姓名检索：沈知夏什么时候入职（对照 WA-104 错字版）。

复用 Anna8 招聘周报 2091834742411366400，不新发报。
用例：WA-140。
"""

from __future__ import annotations

from _paths import CATALOG_PATH, FIXTURES_DIR, ROLE_PERIOD_PEOPLE_PATH, XLSX_PATH
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

CASE_HEADER_ROW = 2
RECORD_HEADER_ROW = 2
RECEIVER = "智本_anrou"
COMPANY = "智本科技"
EXPECT_TOOLS = (
    '{"required":["memory_search"],"forbidden":["memory_settle"],"bound_entity_only":true}'
)

CASE = {
    "用例ID": "WA-140",
    "一级模块": "自然语言选周报",
    "场景": "精确关键词命中",
    "优先级": "P1",
    "优先级说明": "正确姓名检索应命中 Anna8 招聘周报入职信息",
    "测试输入/用户话术": "沈知夏什么时候入职？工号多少？",
    "前置条件": (
        "Anna8 招聘周报（2026/06/15–06/19，reportId=2091834742411366400）"
        "写的是沈知夏、预计 2026 年 9 月 1 日入职、工号 SZ-1907。不附带周报。"
    ),
    "执行步骤": (
        f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
        "2. 不附带周报；\n"
        "3. 用正确姓名「沈知夏」问入职时间与工号；\n"
        "4. 应召回招聘周报并答出 9 月 1 日、SZ-1907。"
    ),
    "必须满足": "沈知夏；9月1日；SZ-1907",
    "禁止出现": "HEALTH-9921；999万；林秋禾；找不到；未命中",
    "预期结果": "精确关键词命中沈知夏招聘周报，答出 9 月 1 日入职、工号 SZ-1907。",
    "评测维度": "忠实度|完整度|指令理解|历史关联",
    "Gold Case": "是",
    "红线用例": "否",
    "建议方式": "自动",
    "需求来源": "向量知识库/精确关键词",
    "expect_tools": EXPECT_TOOLS,
}


def _sheet_headers(ws: Worksheet, header_row: int) -> list[str]:
    return [
        str(cell.value or "").strip()
        for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row))
    ]


def _copy_row_style(ws: Worksheet, src_row: int, dst_row: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)
            dst.alignment = copy(src.alignment)


def write_xlsx(path: Path = XLSX_PATH) -> None:
    wb = load_workbook(path)
    case_ws = wb["用例库"]
    record_ws = wb["执行记录"]
    case_headers = _sheet_headers(case_ws, CASE_HEADER_ROW)
    record_headers = _sheet_headers(record_ws, RECORD_HEADER_ROW)
    case_idx = {name: i + 1 for i, name in enumerate(case_headers) if name}
    rec_case_col = record_headers.index("用例ID") + 1 if "用例ID" in record_headers else None
    existing_rec = set()
    if rec_case_col:
        for row in range(RECORD_HEADER_ROW + 1, record_ws.max_row + 1):
            cid = str(record_ws.cell(row, rec_case_col).value or "").strip()
            if cid:
                existing_rec.add(cid)

    last_case_row = CASE_HEADER_ROW
    case_rows: dict[str, int] = {}
    for row in range(CASE_HEADER_ROW + 1, case_ws.max_row + 1):
        cid = str(case_ws.cell(row, case_idx["用例ID"]).value or "").strip()
        if cid:
            last_case_row = row
            case_rows[cid] = row

    cid = CASE["用例ID"]
    if cid in case_rows:
        target = case_rows[cid]
        print(f"[UPDATE] {cid} {CASE['场景']}")
    else:
        last_case_row += 1
        _copy_row_style(case_ws, last_case_row - 1, last_case_row, len(case_headers))
        target = last_case_row
        case_rows[cid] = target
        print(f"[CASE] {cid} {CASE['场景']}")
    for key, value in CASE.items():
        if key in case_idx:
            case_ws.cell(target, case_idx[key], value=value)
    if rec_case_col and cid not in existing_rec:
        last_rec = RECORD_HEADER_ROW
        for row in range(RECORD_HEADER_ROW + 1, record_ws.max_row + 1):
            if str(record_ws.cell(row, rec_case_col).value or "").strip():
                last_rec = row
        last_rec += 1
        _copy_row_style(record_ws, last_rec - 1, last_rec, len(record_headers))
        record_ws.cell(last_rec, rec_case_col, value=cid)
        existing_rec.add(cid)

    wb.save(path)
    print(f"[XLSX] {cid} -> {path.name}")


def write_catalog(path: Path = CATALOG_PATH) -> None:
    data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    notes = data.setdefault("case_notes", {})
    notes["WA-140"] = (
        "不附带周报，正确姓名「沈知夏」应召回 9月1日入职 SZ-1907（对照 WA-104 错字）"
    )
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("[CATALOG] wrote WA-140")


def main() -> int:
    write_xlsx()
    write_catalog()
    print("[NEXT] python run_report_agent_cases.py --case-ids WA-140 --no-webhook")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
