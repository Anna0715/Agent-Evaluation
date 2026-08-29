#!/usr/bin/env python3
"""补充 1–8 月 ROLE-PERIOD 历史内容检索用例，以及未读周报按人分组总结用例。

复用已发给 智本_anrou 的 ROLE-PERIOD-2026 语料，不新发报。
用例：WA-120 ~ WA-128。
"""

from __future__ import annotations

from _paths import CATALOG_PATH
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from env_config import test_case_xlsx_path

XLSX_PATH = test_case_xlsx_path()
CASE_HEADER_ROW = 2
RECORD_HEADER_ROW = 2
RECEIVER = "智本_anrou"
COMPANY = "智本科技"
DATASET_MARKER = "ROLE-PERIOD-2026"
EXPECT_TOOLS = (
    '{"required":["memory_search"],"forbidden":["memory_settle"],"bound_entity_only":true}'
)

PRECONDITION = (
    f"追问账号={RECEIVER}；追问公司={COMPANY}。"
    f"1–8 月角色周报/月报已发给 {RECEIVER}，标记 {DATASET_MARKER}。"
    "发送人与项目：智本_Anna6(测试勿动)=青检台、智本_Anna8=灰灯回归、"
    "智本_Anna=银杏关账、智本_Anna12=霜桥网关、智本_Anna10=岚图路线图。"
    "不附带周报，由 Agent 检索历史内容。"
)

CASES: list[dict[str, str]] = [
    {
        "用例ID": "WA-120",
        "一级模块": "历史内容检索",
        "场景": "历史月份检索",
        "优先级": "P1",
        "优先级说明": "按月份+项目检索 1–8 月角色周报",
        "测试输入/用户话术": "帮我找一下3月青检台相关的周报",
        "前置条件": PRECONDITION,
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送「帮我找一下3月青检台相关的周报」；\n"
            "4. 应命中 Anna6 青检台 3 月周报要点与探针。"
        ),
        "必须满足": "青检台；QINGJIAN-QA-4406；QA-ANNA6-CANARY-2026；3月",
        "禁止出现": "霜桥网关；RD-ANNA12-CANARY-2026；银杏关账；FIN-ANNA-CANARY-2026",
        "预期结果": "检索到 3 月青检台周报，覆盖项目代号与探针，不串写其它角色专有项目。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "历史内容检索/1-8月",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-121",
        "一级模块": "历史内容检索",
        "场景": "历史月份检索",
        "优先级": "P1",
        "优先级说明": "跨月区间检索财务银杏关账",
        "测试输入/用户话术": "帮我检索一下1月到2月银杏关账的进展",
        "前置条件": PRECONDITION,
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送「帮我检索一下1月到2月银杏关账的进展」；\n"
            "4. 应覆盖财务 Anna 的 1–2 月银杏关账要点。"
        ),
        "必须满足": "银杏关账；YINXING-FIN-2209；FIN-ANNA-CANARY-2026；1月；2月",
        "禁止出现": "青检台；QA-ANNA6-CANARY-2026；霜桥网关；RD-ANNA12-CANARY-2026",
        "预期结果": "覆盖 1–2 月银杏关账进展与探针，不得串写质量/研发专有标记。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "历史内容检索/1-8月",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-122",
        "一级模块": "历史内容检索",
        "场景": "跨月对比检索",
        "优先级": "P1",
        "优先级说明": "上半年与 7–8 月对比同一产品项目",
        "测试输入/用户话术": "对比一下岚图路线图上半年和7到8月的周报重点有什么变化",
        "前置条件": PRECONDITION,
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送对比话术；\n"
            "4. 应点名岚图路线图，并体现上半年与 7–8 月差异或分期要点。"
        ),
        "必须满足": "岚图路线图；LANTU-PM-6615；PM-ANNA10-CANARY-2026；上半年；7月；8月",
        "禁止出现": "银杏关账；FIN-ANNA-CANARY-2026；霜桥网关；RD-ANNA12-CANARY-2026",
        "预期结果": "按上半年与 7–8 月对比岚图路线图周报，命中产品探针，不串财务/研发项目。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "历史内容检索/1-8月",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-123",
        "一级模块": "历史内容检索",
        "场景": "角色历史检索",
        "优先级": "P1",
        "优先级说明": "按角色+项目检索研发历史周报",
        "测试输入/用户话术": "帮我找研发工程师霜桥网关相关的历史周报",
        "前置条件": PRECONDITION,
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送角色历史检索话术；\n"
            "4. 应命中 Anna12 霜桥网关要点与探针。"
        ),
        "必须满足": "霜桥网关；SHUANGQIAO-RD-8830；RD-ANNA12-CANARY-2026；智本_Anna12",
        "禁止出现": "青检台；QA-ANNA6-CANARY-2026；银杏关账；FIN-ANNA-CANARY-2026",
        "预期结果": "检索到研发霜桥网关历史周报，覆盖代号与探针。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "历史内容检索/1-8月",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-124",
        "一级模块": "历史内容检索",
        "场景": "角色历史检索",
        "优先级": "P1",
        "优先级说明": "按产品角色追问 1–8 月关键口径",
        "测试输入/用户话术": "产品角色1到8月周报里岚图路线图的关键口径是多少",
        "前置条件": PRECONDITION,
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送口径追问；\n"
            "4. 应答出岚图路线图专项口径/探针。"
        ),
        "必须满足": "岚图路线图；LANTU-PM-6615；PM-ANNA10-CANARY-2026；1948000",
        "禁止出现": "318600；YINXING-FIN-2209；RD-ANNA12-CANARY-2026",
        "预期结果": "答出产品岚图路线图专项口径 1948000，并带上探针。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "历史内容检索/1-8月",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-125",
        "一级模块": "历史内容检索",
        "场景": "历史月报检索",
        "优先级": "P1",
        "优先级说明": "按人+月份检索月报",
        "测试输入/用户话术": "帮我总结智本_Anna10的5月月报",
        "前置条件": PRECONDITION,
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送「帮我总结智本_Anna10的5月月报」；\n"
            "4. 应基于 5 月月报总结岚图路线图要点。"
        ),
        "必须满足": "智本_Anna10；5月；岚图路线图；PM-ANNA10-CANARY-2026；月报",
        "禁止出现": "FIN-ANNA-CANARY-2026；银杏关账；RD-ANNA12-CANARY-2026",
        "预期结果": "总结 Anna10 5 月月报，覆盖岚图路线图与探针，不串其它角色。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "历史内容检索/1-8月",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-126",
        "一级模块": "历史内容检索",
        "场景": "历史月份检索",
        "优先级": "P1",
        "优先级说明": "7 月质量灰灯回归检索",
        "测试输入/用户话术": "找一下7月灰灯回归相关周报",
        "前置条件": PRECONDITION,
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送「找一下7月灰灯回归相关周报」；\n"
            "4. 应命中 Anna8 7 月灰灯回归要点。"
        ),
        "必须满足": "灰灯回归；GRAYLAMP-QA-7712；QA-ANNA8-CANARY-2026；7月",
        "禁止出现": "岚图路线图；PM-ANNA10-CANARY-2026；银杏关账；FIN-ANNA-CANARY-2026",
        "预期结果": "检索到 7 月灰灯回归周报，覆盖代号与探针。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "历史内容检索/1-8月",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-127",
        "一级模块": "历史内容检索",
        "场景": "跨月对比检索",
        "优先级": "P1",
        "优先级说明": "同一项目 3 月与 8 月对比",
        "测试输入/用户话术": "青检台3月和8月的周报分别做了什么",
        "前置条件": PRECONDITION,
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送跨月对比话术；\n"
            "4. 应分别覆盖 3 月与 8 月青检台要点。"
        ),
        "必须满足": "青检台；QINGJIAN-QA-4406；QA-ANNA6-CANARY-2026；3月；8月",
        "禁止出现": "霜桥网关；RD-ANNA12-CANARY-2026；岚图路线图；PM-ANNA10-CANARY-2026",
        "预期结果": "分别说明青检台 3 月与 8 月周报要点，命中探针，不串其它项目。",
        "评测维度": "忠实度|完整度|指令理解|历史关联",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动",
        "需求来源": "历史内容检索/1-8月",
        "expect_tools": EXPECT_TOOLS,
    },
    {
        "用例ID": "WA-128",
        "一级模块": "收件箱未读",
        "场景": "未读周报按人分组",
        "优先级": "P0",
        "优先级说明": "统计未读周报并按汇报人分组后分别总结",
        "测试输入/用户话术": (
            "帮我统计下我目前未读的周报数据，并按照汇报人分组，再分别进行未读的汇报总结"
        ),
        "前置条件": (
            f"追问账号={RECEIVER}；追问公司={COMPANY}。"
            "执行时按收件箱 scope=received、category=周报、readStatus=未读 实时拉取；"
            "不附带周报，由 Agent 自行统计未读并按汇报人分组总结。"
            f"若未读中含 {DATASET_MARKER} 语料，应按发送人分组覆盖其项目要点。"
        ),
        "执行步骤": (
            f"1. 以 {RECEIVER} / {COMPANY} 登录；\n"
            "2. 不附带周报；\n"
            "3. 发送未读统计+按人分组总结话术；\n"
            "4. 应先统计未读，再按汇报人分组并分别总结；无未读时明确提示。"
        ),
        "必须满足": (
            "按汇报人分组或等价分人结构；覆盖当前未读周报中的汇报人；"
            "对每位汇报人分别总结未读周报要点；无未读时明确提示暂无未读。"
        ),
        "禁止出现": "不得虚构未读周报；不得把已读周报或日报/月报当成未读周报；不得把进行中写成已完成。",
        "预期结果": (
            "先给出未读周报统计，再按汇报人分组并分别总结；"
            "若当前无未读周报，明确提示没有未读。"
        ),
        "评测维度": "指令理解|完整度|忠实度|可读性",
        "Gold Case": "是",
        "红线用例": "否",
        "建议方式": "自动+人工",
        "需求来源": "未读周报按人分组总结",
        "expect_tools": EXPECT_TOOLS,
    },
]


def _sheet_headers(ws: Worksheet, header_row: int) -> list[str]:
    return [
        str(cell.value or "").strip()
        for cell in next(ws.iter_rows(min_row=header_row, max_row=header_row))
    ]


def _copy_row_style(ws: Worksheet, src_row: int, dst_row: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)
            dst.alignment = copy(src.alignment)


def write_xlsx(path: Path = XLSX_PATH) -> None:
    wb = load_workbook(path)
    case_ws = wb["用例库"]
    record_ws = wb["执行记录"]
    case_headers = _sheet_headers(case_ws, CASE_HEADER_ROW)
    record_headers = _sheet_headers(record_ws, RECORD_HEADER_ROW)
    case_idx = {name: i + 1 for i, name in enumerate(case_headers) if name}
    rec_case_col = record_headers.index("用例ID") + 1 if "用例ID" in record_headers else None
    existing_rec = set()
    if rec_case_col:
        for row in range(RECORD_HEADER_ROW + 1, record_ws.max_row + 1):
            cid = str(record_ws.cell(row, rec_case_col).value or "").strip()
            if cid:
                existing_rec.add(cid)

    last_case_row = CASE_HEADER_ROW
    case_rows: dict[str, int] = {}
    for row in range(CASE_HEADER_ROW + 1, case_ws.max_row + 1):
        cid = str(case_ws.cell(row, case_idx["用例ID"]).value or "").strip()
        if cid:
            last_case_row = row
            case_rows[cid] = row

    for spec in CASES:
        cid = spec["用例ID"]
        if cid in case_rows:
            target = case_rows[cid]
            print(f"[UPDATE] {cid} {spec['场景']}")
        else:
            last_case_row += 1
            _copy_row_style(case_ws, last_case_row - 1, last_case_row, len(case_headers))
            target = last_case_row
            case_rows[cid] = target
            print(f"[CASE] {cid} {spec['场景']}")
        for key, value in spec.items():
            if key in case_idx:
                case_ws.cell(target, case_idx[key], value=value)
        if rec_case_col and cid not in existing_rec:
            last_rec = RECORD_HEADER_ROW
            for row in range(RECORD_HEADER_ROW + 1, record_ws.max_row + 1):
                if str(record_ws.cell(row, rec_case_col).value or "").strip():
                    last_rec = row
            last_rec += 1
            _copy_row_style(record_ws, last_rec - 1, last_rec, len(record_headers))
            record_ws.cell(last_rec, rec_case_col, value=cid)
            existing_rec.add(cid)

    wb.save(path)
    print(f"[XLSX] WA-120..128 -> {path.name}")


def write_catalog(path: Path = CATALOG_PATH) -> None:
    data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    notes = data.setdefault("case_notes", {})
    notes.update(
        {
            "WA-120": "不附带周报，检索 3 月青检台 ROLE-PERIOD 周报",
            "WA-121": "不附带周报，检索 1–2 月银杏关账进展",
            "WA-122": "不附带周报，对比岚图路线图上半年与 7–8 月",
            "WA-123": "不附带周报，检索研发霜桥网关历史周报",
            "WA-124": "不附带周报，追问产品岚图路线图 1–8 月口径",
            "WA-125": "不附带周报，总结 Anna10 5 月月报",
            "WA-126": "不附带周报，检索 7 月灰灯回归周报",
            "WA-127": "不附带周报，对比青检台 3 月与 8 月周报",
            "WA-128": "不附带周报，统计未读周报并按汇报人分组分别总结",
        }
    )
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("[CATALOG] wrote WA-120..128")


def main() -> int:
    write_xlsx()
    write_catalog()
    ids = ",".join(item["用例ID"] for item in CASES)
    print(f"[NEXT] python run_report_agent_cases.py --case-ids {ids}")
    print("       增量 webhook 仅失败时发送（全部通过跳过；--no-webhook 可关闭）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
