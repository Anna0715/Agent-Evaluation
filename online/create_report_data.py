#!/usr/bin/env python3
"""基于周报 Agent 测试集 / report_data.csv，批量构造日报/周报/月报数据。

创建流程（与客户端一致）：
1. POST /reports/create（status=0，不传 reportId）拿到服务端分配的 reportId
2. 再 POST /reports/create，带上该 reportId、正文、汇报对象、startDate、status=1

种子造数：
    python create_report_data.py
    python create_report_data.py --execute --count 6

从 report_data.csv 发报并轮询 AI 摘要（send_and_poll_weekly_report.sh 调用）：
    python create_report_data.py --from-csv report_data.csv --execute --poll
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import os
import re
import subprocess
import sys
import time
import uuid
from calendar import monthrange
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

SCRIPT_DIR = Path(__file__).resolve().parent
LOGIN_SCRIPT = SCRIPT_DIR / "login.py"
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from env_config import FIXTURES_DIR  # noqa: E402

from excel_tool import DATA_XLSX_PATH, Account, get_account_from_data_csv, get_accounts  # noqa: E402
from env_config import (  # noqa: E402
    apply_env,
    dataset_version_label,
    report_data_csv,
    report_data_dir,
    report_data_xlsx,
    reports_create_url,
    reports_detail_url,
    saas_api_url,
)

REPORTS_CREATE_URL = reports_create_url()
REPORTS_DETAIL_URL = reports_detail_url()
DEFAULT_LOGIN_BASE_URL = saas_api_url()
DEFAULT_COMPANY_NAME = "智本科技"
DEFAULT_USER_NAME = "智本_anrou"
DEFAULT_REPORT_TO_USER_ID = "ouvfgsszpinfjg"
DEFAULT_REPORT_TO_NAME = "公台"
DEFAULT_REPORT_TO_AVATAR = (
    "https://test-saas-api.oa-test.org/api/oa/v1/file/fl09TqR8RUxXUg3i0VFgssjalskce"
)
DEFAULT_REPORT_CSV = report_data_csv()
DEFAULT_REPORT_CSV_DIR = report_data_dir()
DEFAULT_REPORT_XLSX = report_data_xlsx()


def refresh_env_config(env: str | None = None) -> None:
    global REPORTS_CREATE_URL, REPORTS_DETAIL_URL, DEFAULT_LOGIN_BASE_URL
    global DEFAULT_REPORT_CSV, DEFAULT_REPORT_CSV_DIR, DEFAULT_REPORT_XLSX
    apply_env(env)
    REPORTS_CREATE_URL = reports_create_url(env)
    REPORTS_DETAIL_URL = reports_detail_url(env)
    DEFAULT_LOGIN_BASE_URL = saas_api_url(env)
    DEFAULT_REPORT_CSV = report_data_csv(env)
    DEFAULT_REPORT_CSV_DIR = report_data_dir(env)
    DEFAULT_REPORT_XLSX = report_data_xlsx(env)
REPORT_CSV_COLUMNS = [
    "日期",
    "发送人",
    "接收人",
    "周报类型",
    "周报内容",
    "AI总结内容",
    "reportId",
]
TOKYO = ZoneInfo("Asia/Tokyo")

CATEGORY_DAILY = 1
CATEGORY_WEEKLY = 2
CATEGORY_MONTHLY = 3
CATEGORY_NAME = {
    CATEGORY_DAILY: "日报",
    CATEGORY_WEEKLY: "周报",
    CATEGORY_MONTHLY: "月报",
}
CATEGORY_BY_NAME = {name: code for code, name in CATEGORY_NAME.items()}
CSV_RESULT_REPORT_ID = "reportId"
CSV_RESULT_AI_SUMMARY = "AI总结内容"

TEMPLATE_HEADERS = {
    CATEGORY_DAILY: ("今日工作：", "明日工作：", "需要协调和帮助："),
    CATEGORY_WEEKLY: ("本周工作：", "下周工作：", "需要协调和帮助："),
    CATEGORY_MONTHLY: ("本月工作：", "下月工作：", "需要协调和帮助："),
}

# 覆盖三种类型；内容与测试集主题对齐，便于后续 Agent 评测对照。
SEED_REPORTS: list[dict[str, Any]] = [
    {
        "category": CATEGORY_DAILY,
        "case_id": "WR-001",
        "offset_days": 0,
        "summary": "完成登录页错误提示优化，修复 BUG-231 并通过本地测试。",
        "progress": "联调登录异常提示文案；本地回归通过。",
        "risk": "线上回归尚未覆盖。",
        "next_plan": "明天补充回归测试。",
    },
    {
        "category": CATEGORY_DAILY,
        "case_id": "WR-013",
        "offset_days": -1,
        "summary": "完成支付接口联调，并修复回调签名问题。",
        "progress": "上午联调支付接口；下午修复回调签名后再次联调通过。",
        "risk": "无阻塞。",
        "next_plan": "补充回调异常场景用例。",
    },
    {
        "category": CATEGORY_WEEKLY,
        "case_id": "WR-002",
        "offset_weeks": 0,
        "summary": "本周完成需求评审，接口联调进度 80%。",
        "progress": "周一完成需求评审；周二接口联调完成 80%；周三识别第三方接口限流风险。",
        "risk": "第三方接口存在限流风险。",
        "next_plan": "继续联调并补充限流兜底方案。",
    },
    {
        "category": CATEGORY_WEEKLY,
        "case_id": "WR-008",
        "offset_weeks": -1,
        "summary": "支付成功率从 97.1% 提升到 98.4%。",
        "progress": "观察支付成功率改善；重试策略优化仍在归因分析中。",
        "risk": "提升原因尚未完成归因，不能断言完全由重试策略导致。",
        "next_plan": "完成归因分析并沉淀监控看板。",
    },
    {
        "category": CATEGORY_WEEKLY,
        "case_id": "WR-011",
        "offset_weeks": -2,
        "summary": "定位并修复缓存命中率下降问题。",
        "progress": "排查缓存命中率下降；定位为缓存键版本不一致；修复完成待发布。",
        "risk": "修复尚未发布到生产。",
        "next_plan": "安排发布并观察命中率回升。",
    },
    {
        "category": CATEGORY_MONTHLY,
        "case_id": "WR-003",
        "offset_months": 0,
        "summary": "7 月新增客户 12 家，续费客户 8 家。",
        "progress": "推进获客与续费；月底数据尚未关闭。",
        "risk": "月末数据未关闭，暂不能视为最终确认值。",
        "next_plan": "月初关闭账期并输出最终月报。",
    },
    {
        "category": CATEGORY_MONTHLY,
        "case_id": "WR-012",
        "offset_months": -1,
        "summary": "活跃用户由 1000 增至 1200，总体增长 20%。",
        "progress": "第 1-4 周活跃用户分别为 1000/1100/1050/1200；第 3 周短暂回落。",
        "risk": "回落原因尚未验证。",
        "next_plan": "分析第 3 周回落原因，巩固增长策略。",
    },
    {
        "category": CATEGORY_MONTHLY,
        "case_id": "WR-024",
        "offset_months": 0,
        "summary": "7 月完成项目启动、一期开发与验收提交。",
        "progress": "7/1 项目启动；7/15 完成一期开发；7/31 提交验收。",
        "risk": "验收结果落在次月，本月不宣称已验收通过。",
        "next_plan": "跟进验收结论并规划二期。",
    },
]


def load_login_module() -> Any:
    if not LOGIN_SCRIPT.exists():
        raise FileNotFoundError(f"登录脚本不存在：{LOGIN_SCRIPT}")
    login_dir = str(LOGIN_SCRIPT.parent)
    if login_dir not in sys.path:
        sys.path.insert(0, login_dir)
    spec = importlib.util.spec_from_file_location("autotest_login", LOGIN_SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法加载登录脚本：{LOGIN_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def monday_of_week(day: date) -> date:
    return day - timedelta(days=day.weekday())


def add_months(day: date, months: int) -> date:
    month = day.month - 1 + months
    year = day.year + month // 12
    month = month % 12 + 1
    return date(year, month, 1)


def start_of_day_ms(day: date) -> int:
    value = datetime(day.year, day.month, day.day, tzinfo=TOKYO)
    return int(value.timestamp() * 1000)


def resolve_period(item: dict[str, Any], today: date) -> tuple[date, date, str]:
    category = int(item["category"])
    if category == CATEGORY_DAILY:
        day = today + timedelta(days=int(item.get("offset_days", 0)))
        label = day.isoformat()
        return day, day, label
    if category == CATEGORY_WEEKLY:
        week_start = monday_of_week(today) + timedelta(weeks=int(item.get("offset_weeks", 0)))
        week_end = week_start + timedelta(days=6)
        label = f"{week_start.isoformat()} 至 {week_end.isoformat()}"
        return week_start, week_end, label
    if category == CATEGORY_MONTHLY:
        month_start = add_months(date(today.year, today.month, 1), int(item.get("offset_months", 0)))
        last_day = monthrange(month_start.year, month_start.month)[1]
        month_end = date(month_start.year, month_start.month, last_day)
        label = f"{month_start.year}-{month_start.month:02d}"
        return month_start, month_end, label
    raise ValueError(f"unsupported category={category}")


def bold_paragraph(text: str) -> dict[str, Any]:
    return {
        "type": "paragraph",
        "attrs": {"indent": 0, "textAlign": None},
        "content": [{"type": "text", "text": text, "marks": [{"type": "bold"}]}],
    }


def text_paragraph(text: str) -> dict[str, Any]:
    return {
        "type": "paragraph",
        "attrs": {"indent": 0, "textAlign": None},
        "content": [{"type": "text", "text": text}],
    }


def empty_paragraph() -> dict[str, Any]:
    return {"type": "paragraph", "attrs": {"indent": 0, "textAlign": None}}


def build_draft_content(category: int) -> dict[str, Any]:
    """客户端首次创建草稿时的空模板。"""
    current, nxt, help_header = TEMPLATE_HEADERS[category]
    return {
        "type": "doc",
        "content": [
            bold_paragraph(current),
            empty_paragraph(),
            bold_paragraph(nxt),
            empty_paragraph(),
            bold_paragraph(help_header),
            empty_paragraph(),
        ],
    }


def build_content(item: dict[str, Any], period_label: str) -> dict[str, Any]:
    """提交时写入的完整正文，结构对齐客户端模板。"""
    category = int(item["category"])
    current, nxt, help_header = TEMPLATE_HEADERS[category]
    name = CATEGORY_NAME[category]
    current_body = (
        f"{name}（{period_label}）\n"
        f"概览：{item['summary']}\n"
        f"进展：{item['progress']}\n"
        f"风险：{item['risk']}"
    )
    next_body = item["next_plan"]
    help_body = f"对照用例：{item.get('case_id', '-')}"
    return {
        "type": "doc",
        "content": [
            bold_paragraph(current),
            text_paragraph(current_body),
            bold_paragraph(nxt),
            text_paragraph(next_body),
            bold_paragraph(help_header),
            text_paragraph(help_body),
        ],
    }


def build_headers(
    *,
    token: str,
    tenant_id: str,
    company_id: str,
    user_id: str,
) -> dict[str, str]:
    request_id = str(uuid.uuid4())
    return {
        "accept": "*/*",
        "accept-language": "zh-CN",
        "authorization": f"Bearer {token}",
        "client-version": "1.0.335",
        "company-id": company_id,
        "content-type": "application/json",
        "device-id": "8CC88A47-6FAE-5F8B-A3A5-4E4536614029",
        "device-name": "Anna's%20%20macbook%20(MacBook%20Pro)",
        "device-os": "macOS 26.5.1",
        "operationid": request_id,
        "priority": "u=1, i",
        "tenant-id": tenant_id,
        "timezone": "9",
        "token": token,
        "user-agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "iHive/1.0.335 Chrome/142.0.7444.265 Electron/39.3.0 Safari/537.36"
        ),
        "user-id": user_id,
        "x-device-arch": "arm64",
        "x-iana-timezone": "Asia/Tokyo",
        "x-request-id": request_id,
        "x-timezone": "9",
    }


def post_json(url: str, headers: dict[str, str], payload: dict[str, Any]) -> dict[str, Any]:
    # 每次请求换新 operationid / x-request-id，避免串单。
    fresh = dict(headers)
    request_id = str(uuid.uuid4())
    fresh["operationid"] = request_id
    fresh["x-request-id"] = request_id

    command = [
        "curl",
        "-sS",
        "--fail-with-body",
        "--max-time",
        "30",
        "-X",
        "POST",
        url,
    ]
    for key, value in fresh.items():
        command.extend(["-H", f"{key}: {value}"])
    command.extend(
        ["--data-raw", json.dumps(payload, ensure_ascii=False, separators=(",", ":"))]
    )
    result = subprocess.run(command, capture_output=True, text=True, check=False)
    body = (result.stdout or "").strip()
    if result.returncode != 0:
        raise RuntimeError(
            f"请求失败：curl={result.returncode} url={url}，"
            f"response={body[:500]!r}，stderr={(result.stderr or '').strip()[:300]!r}"
        )
    if not body:
        return {"success": True, "raw": ""}
    try:
        return json.loads(body)
    except json.JSONDecodeError:
        return {"success": True, "raw": body[:500]}


def post_report(headers: dict[str, str], payload: dict[str, Any]) -> dict[str, Any]:
    return post_json(REPORTS_CREATE_URL, headers, payload)


def response_ok(response: dict[str, Any]) -> bool:
    err_code = response.get("errCode", response.get("errcode", 0))
    return not (isinstance(err_code, int) and err_code != 0)


def extract_report_id(response: dict[str, Any]) -> str:
    data = response.get("data")
    if isinstance(data, dict):
        for key in ("reportId", "report_id", "id"):
            value = data.get(key)
            if isinstance(value, (str, int)) and str(value).strip():
                return str(value).strip()
        report = data.get("report")
        if isinstance(report, dict):
            for key in ("reportId", "report_id", "id"):
                value = report.get(key)
                if isinstance(value, (str, int)) and str(value).strip():
                    return str(value).strip()
    for key in ("reportId", "report_id", "id"):
        value = response.get(key)
        if isinstance(value, (str, int)) and str(value).strip():
            return str(value).strip()
    return ""


def allocate_report_id(headers: dict[str, str], category: int) -> tuple[str, dict[str, Any]]:
    """第一步：创建草稿，从返回中取 reportId。"""
    payload = {
        "category": category,
        "content": build_draft_content(category),
        "status": 0,
    }
    response = post_report(headers, payload)
    if not response_ok(response):
        raise RuntimeError(f"分配 reportId 失败：{response}")
    report_id = extract_report_id(response)
    if not report_id:
        raise RuntimeError(f"草稿创建成功但未返回 reportId：{response}")
    return report_id, response


def select_seeds(categories: set[int], count: int) -> list[dict[str, Any]]:
    filtered = [item for item in SEED_REPORTS if int(item["category"]) in categories]
    if not filtered:
        raise ValueError("没有匹配的种子数据")
    if count > len(filtered):
        raise ValueError(f"--count 超过可用种子数 {len(filtered)}")
    return filtered[:count]


def resolve_auth(args: argparse.Namespace) -> dict[str, str]:
    if args.auth_from_env:
        token = os.environ.get("REPORT_TOKEN", "").strip()
        tenant_id = os.environ.get("REPORT_TENANT_ID", "").strip()
        company_id = os.environ.get("REPORT_COMPANY_ID", "").strip()
        user_id = os.environ.get("REPORT_USER_ID", "").strip()
        missing = [
            name
            for name, value in (
                ("REPORT_TOKEN", token),
                ("REPORT_TENANT_ID", tenant_id),
                ("REPORT_COMPANY_ID", company_id),
                ("REPORT_USER_ID", user_id),
            )
            if not value
        ]
        if missing:
            raise RuntimeError(f"--auth-from-env 缺少环境变量：{', '.join(missing)}")
        return {
            "token": token,
            "tenant_id": tenant_id,
            "company_id": company_id,
            "user_id": user_id,
        }

    login_module = load_login_module()
    login_module.ORG_BASE_URL = args.login_base_url.rstrip("/")
    login_result = login_module.login(
        company_name=args.company_name,
        user_name=args.user_name,
    )
    required = ("tenant_id", "company_id")
    missing = [key for key in required if not login_result.get(key)]
    if missing:
        raise RuntimeError(f"登录结果缺少字段：{', '.join(missing)}")
    user_id = login_result.get("user_id") or login_result.get("userid")
    if not user_id:
        raise RuntimeError("登录结果缺少 user_id/userid")
    token = login_result.get("api_token") or login_result.get("im_token") or login_result.get(
        "token_im_access_token"
    )
    if not token:
        raise RuntimeError("登录结果缺少 api_token/im_token/token_im_access_token")
    return {
        "token": token,
        "tenant_id": login_result["tenant_id"],
        "company_id": login_result["company_id"],
        "user_id": user_id,
    }


def build_content_from_text(text: str) -> dict[str, Any]:
    """把 CSV「周报内容」转成客户端 doc 结构。"""
    lines = (text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    paragraphs: list[dict[str, Any]] = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            paragraphs.append(empty_paragraph())
        else:
            paragraphs.append(text_paragraph(stripped))
    if not paragraphs:
        paragraphs.append(empty_paragraph())
    return {"type": "doc", "content": paragraphs}


def parse_single_day(value: str) -> date:
    text = value.strip().replace(".", "-").replace("/", "-")
    return datetime.strptime(text, "%Y-%m-%d").date()


def parse_csv_period(date_text: str, category: int) -> tuple[date, date, str]:
    """按「日期 + 周报类型」解析周期；周报按表格中的起止范围。"""
    raw = (date_text or "").strip()
    if not raw:
        raise ValueError("日期为空")

    if category == CATEGORY_DAILY:
        day = parse_single_day(raw)
        return day, day, day.isoformat()

    if category == CATEGORY_WEEKLY:
        # 2026/07/01-07/03 或 2026-07-01~2026-07-03 或 2026/07/01至07/03
        matched = re.match(
            r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2})\s*[-~～至]+\s*(?:(\d{4})[/-])?(\d{1,2})[/-](\d{1,2})$",
            raw,
        )
        if not matched:
            # 兜底：单日则按该日所在周（周一到周日）
            day = parse_single_day(raw)
            week_start = monday_of_week(day)
            week_end = week_start + timedelta(days=6)
            label = f"{week_start.isoformat()} 至 {week_end.isoformat()}"
            return week_start, week_end, label
        year = int(matched.group(1))
        start = date(year, int(matched.group(2)), int(matched.group(3)))
        end_year = int(matched.group(4) or year)
        end = date(end_year, int(matched.group(5)), int(matched.group(6)))
        label = f"{start.isoformat()} 至 {end.isoformat()}"
        return start, end, label

    if category == CATEGORY_MONTHLY:
        matched = re.match(r"^(\d{4})\s*年\s*(\d{1,2})\s*月$", raw)
        if matched:
            year, month = int(matched.group(1)), int(matched.group(2))
        else:
            matched = re.match(r"^(\d{4})[/-](\d{1,2})$", raw)
            if not matched:
                raise ValueError(f"无法解析月报日期：{raw!r}")
            year, month = int(matched.group(1)), int(matched.group(2))
        month_start = date(year, month, 1)
        last_day = monthrange(year, month)[1]
        month_end = date(year, month, last_day)
        return month_start, month_end, f"{year}-{month:02d}"

    raise ValueError(f"unsupported category={category}")


def lookup_account(user_name: str, company_name: str = "") -> Account:
    import excel_tool

    xlsx_path = excel_tool.DATA_XLSX_PATH
    name = (user_name or "").strip()
    if not name:
        raise ValueError("账号名为空")
    company = (company_name or "").strip()
    if company:
        return get_account_from_data_csv(xlsx_path, user_name=name, company_name=company)
    matches = [account for account in get_accounts(xlsx_path) if account.user_name == name]
    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ValueError(f"data.xlsx 中找不到 UserName={name!r}")
    companies = sorted({item.company_name for item in matches})
    raise ValueError(f"data.xlsx 中 UserName={name!r} 匹配多条，请指定公司：{companies}")


def login_as_user(
    *,
    user_name: str,
    company_name: str,
    login_base_url: str,
    cache: dict[tuple[str, str], dict[str, str]],
) -> dict[str, str]:
    key = (user_name, company_name)
    cached = cache.get(key)
    if cached:
        return cached

    account = lookup_account(user_name, company_name)
    resolved_company = company_name or account.company_name or DEFAULT_COMPANY_NAME
    login_module = load_login_module()
    login_module.ORG_BASE_URL = login_base_url.rstrip("/")
    login_result = login_module.login(company_name=resolved_company, user_name=user_name)
    user_id = login_result.get("user_id") or login_result.get("userid") or account.user_id
    if not user_id:
        raise RuntimeError(f"登录结果缺少 user_id：{user_name}")
    # 与种子模式 resolve_auth 一致：优先 api_token。
    token = (
        login_result.get("api_token")
        or login_result.get("im_token")
        or login_result.get("token_im_access_token")
    )
    if not token:
        raise RuntimeError(f"登录结果缺少 token：{user_name}")
    auth = {
        "token": token,
        "tenant_id": login_result["tenant_id"],
        "company_id": login_result["company_id"],
        "user_id": str(user_id),
        "user_name": user_name,
        "company_name": resolved_company,
    }
    cache[key] = auth
    return auth


def format_ai_summary(ai_summary: Any) -> str:
    if not isinstance(ai_summary, dict):
        return "" if ai_summary is None else str(ai_summary)
    parts: list[str] = []
    summary = ai_summary.get("summary")
    if summary:
        parts.append(str(summary).strip())
    content = ai_summary.get("content") or []
    if isinstance(content, list):
        for block in content:
            if not isinstance(block, dict):
                continue
            title = str(block.get("title") or "").strip()
            if title:
                parts.append(f"【{title}】")
            for item in block.get("items") or []:
                if not isinstance(item, dict):
                    continue
                status = str(item.get("status") or "").strip()
                text = str(item.get("text") or "").strip()
                due = item.get("due")
                due_s = f" (截止:{due})" if due else ""
                status_s = f"[{status}]" if status else ""
                parts.append(f"  {status_s}{due_s} {text}".rstrip())
    return "\n".join(part for part in parts if part).strip()


def print_ai_summary_like_original(ai_summary: Any) -> None:
    """对齐原 send_and_poll_weekly_report.sh 的摘要打印格式。"""
    if not isinstance(ai_summary, dict):
        print(ai_summary)
        return
    print("----- aiSummary.summary -----")
    print(ai_summary.get("summary") or "")
    print("----- aiSummary 明细 -----")
    content = ai_summary.get("content") or []
    if isinstance(content, list):
        for block in content:
            if not isinstance(block, dict):
                continue
            title = block.get("title") or ""
            print(f"【{title}】")
            for item in block.get("items") or []:
                if not isinstance(item, dict):
                    continue
                status = item.get("status") or ""
                text = item.get("text") or ""
                due = item.get("due")
                due_s = f" (截止:{due})" if due else ""
                print(f"  [{status}]{due_s} {text}")
    print("----- aiSummary 原始 JSON -----")
    print(json.dumps(ai_summary, ensure_ascii=False, indent=2))


def poll_ai_summary(
    *,
    receiver_auth: dict[str, str],
    report_id: str,
    poll_interval: float,
    max_polls: int,
) -> tuple[str, dict[str, Any], Any]:
    """接收人轮询 /reports/detail，直到 aiTaskStatus=succeeded。

    返回：(格式化文本, detail 响应, aiSummary 原始对象)
    """
    headers = build_headers(
        token=receiver_auth["token"],
        tenant_id=receiver_auth["tenant_id"],
        company_id=receiver_auth["company_id"],
        user_id=receiver_auth["user_id"],
    )
    body = {"reportId": report_id, "tenantId": receiver_auth["tenant_id"]}
    last: dict[str, Any] = {}
    for index in range(1, max_polls + 1):
        last = post_json(REPORTS_DETAIL_URL, headers, body)
        report = ((last.get("data") or {}) if isinstance(last.get("data"), dict) else {}).get("report")
        report = report if isinstance(report, dict) else {}
        status = str(report.get("aiTaskStatus") or "unknown")
        err_code = last.get("errCode", last.get("errcode", "?"))
        print(f"    #{index} errCode={err_code} aiTaskStatus={status}", flush=True)
        if status == "succeeded":
            ai_summary = report.get("aiSummary")
            return format_ai_summary(ai_summary), last, ai_summary
        if status == "failed":
            raise RuntimeError(f"摘要任务失败 reportId={report_id} response={last}")
        if index < max_polls:
            time.sleep(poll_interval)
    raise TimeoutError(
        f"超时: {max_polls} 次轮询后 aiTaskStatus 仍未 succeeded，reportId={report_id}"
    )


def load_report_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    path = Path(path)
    if path.is_dir():
        return load_report_csv_dir(path)
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError(f"CSV 无表头：{path}")
        fieldnames = list(reader.fieldnames)
        rows = [{key: (row.get(key) or "") for key in fieldnames} for row in reader]
    for col in (CSV_RESULT_REPORT_ID, CSV_RESULT_AI_SUMMARY):
        if col not in fieldnames:
            fieldnames.append(col)
            for row in rows:
                row[col] = ""
    required = {"日期", "发送人", "接收人", "周报类型", "周报内容"}
    missing = required - set(fieldnames)
    if missing:
        raise ValueError(f"CSV 缺少列：{sorted(missing)}")
    return fieldnames, rows


def _sender_csv_filename(sender: str) -> str:
    safe = re.sub(r'[<>:"/\\|?*]', "_", (sender or "未命名").strip()) or "未命名"
    return f"{safe}.csv"


def _excel_sheet_name(sender: str, used: set[str]) -> str:
    """Excel sheet 名 = 发送人名称（去掉非法字符，最长 31）。"""
    raw = re.sub(r"[\[\]:*?/\\]", "_", (sender or "未命名").strip()) or "未命名"
    name = raw[:31]
    base = name
    suffix = 1
    while name in used:
        tail = f"_{suffix}"
        name = f"{base[: 31 - len(tail)]}{tail}"
        suffix += 1
    used.add(name)
    return name


def _group_rows_by_sender(rows: list[dict[str, str]]) -> tuple[list[str], dict[str, list[dict[str, str]]]]:
    grouped: dict[str, list[dict[str, str]]] = {}
    sender_order: list[str] = []
    for row in rows:
        sender = (row.get("发送人") or "").strip() or "未命名"
        if sender not in grouped:
            sender_order.append(sender)
            grouped[sender] = []
        grouped[sender].append(row)
    return sender_order, grouped


def export_report_data_xlsx(
    csv_path: Path | str = DEFAULT_REPORT_CSV,
    xlsx_path: Path | str = DEFAULT_REPORT_XLSX,
) -> int:
    """从 report_data.csv 生成 report_data.xlsx：每个发送人一个 sheet，sheet 名=发送人。"""
    from openpyxl import Workbook

    csv_path = Path(csv_path)
    xlsx_path = Path(xlsx_path)
    fieldnames, rows = load_report_csv(csv_path)
    columns = [col for col in REPORT_CSV_COLUMNS if col in fieldnames] or list(fieldnames)
    sender_order, grouped = _group_rows_by_sender(rows)

    wb = Workbook()
    default = wb.active
    used_names: set[str] = set()
    first = True
    for sender in sender_order:
        sheet_name = _excel_sheet_name(sender, used_names)
        ws = default if first else wb.create_sheet(sheet_name)
        if first:
            ws.title = sheet_name
            first = False
        for col_idx, header in enumerate(columns, start=1):
            ws.cell(1, col_idx, value=header)
        for row_idx, row in enumerate(grouped[sender], start=2):
            for col_idx, header in enumerate(columns, start=1):
                ws.cell(row_idx, col_idx, value=row.get(header, ""))

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(
        f"[XLSX] sheets={len(sender_order)} rows={len(rows)} "
        f"sheet名=发送人 from={csv_path.name} -> {xlsx_path.name}",
        flush=True,
    )
    for sender in sender_order:
        print(f"  sheet「{sender}」: {len(grouped[sender])} 行", flush=True)
    return len(rows)


def load_report_csv_dir(dir_path: Path) -> tuple[list[str], list[dict[str, str]]]:
    dir_path = Path(dir_path)
    files = sorted(dir_path.glob("*.csv"))
    if not files:
        raise FileNotFoundError(f"目录下无 CSV：{dir_path}")
    fieldnames: list[str] = []
    rows: list[dict[str, str]] = []
    for file_path in files:
        names, part = load_report_csv(file_path)
        if not fieldnames:
            fieldnames = names
        rows.extend(part)
    return fieldnames, rows


def split_report_csv_by_sender(
    csv_path: Path | str = DEFAULT_REPORT_CSV,
    out_dir: Path | str = DEFAULT_REPORT_CSV_DIR,
    *,
    also_xlsx: bool = True,
    xlsx_path: Path | str = DEFAULT_REPORT_XLSX,
) -> int:
    """把 report_data.csv 按发送人拆成 report_data/{发送人}.csv，并可选生成多 sheet 的 xlsx。"""
    csv_path = Path(csv_path)
    out_dir = Path(out_dir)
    fieldnames, rows = load_report_csv(csv_path)
    sender_order, grouped = _group_rows_by_sender(rows)

    out_dir.mkdir(parents=True, exist_ok=True)
    for stale in out_dir.glob("*.csv"):
        stale.unlink()

    total = 0
    for sender in sender_order:
        part = grouped[sender]
        write_report_csv(out_dir / _sender_csv_filename(sender), fieldnames, part)
        total += len(part)
        print(f"  {sender}: {len(part)} -> {out_dir.name}/{_sender_csv_filename(sender)}")

    print(
        f"[SPLIT] senders={len(sender_order)} rows={total} "
        f"from={csv_path.name} -> {out_dir}/",
        flush=True,
    )
    if also_xlsx:
        export_report_data_xlsx(csv_path, xlsx_path)
    return total


def merge_report_csv_dir(
    dir_path: Path | str = DEFAULT_REPORT_CSV_DIR,
    out_path: Path | str = DEFAULT_REPORT_CSV,
) -> int:
    """把 report_data/*.csv 合并回 report_data.csv。"""
    dir_path = Path(dir_path)
    out_path = Path(out_path)
    fieldnames, rows = load_report_csv_dir(dir_path)
    write_report_csv(out_path, fieldnames, rows)
    print(
        f"[MERGE] files={len(list(dir_path.glob('*.csv')))} rows={len(rows)} "
        f"-> {out_path.name}",
        flush=True,
    )
    return len(rows)


def write_report_csv(path: Path, fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key, "") for key in fieldnames})


def filter_csv_rows(
    rows: list[dict[str, str]],
    *,
    type_filter: set[str],
    limit: int,
    force: bool,
    new_only: bool = False,
) -> list[tuple[int, dict[str, str]]]:
    selected: list[tuple[int, dict[str, str]]] = []
    for index, row in enumerate(rows):
        type_name = (row.get("周报类型") or "").strip()
        if type_filter and type_name not in type_filter:
            continue
        has_report_id = bool((row.get(CSV_RESULT_REPORT_ID) or "").strip())
        has_ai = bool((row.get(CSV_RESULT_AI_SUMMARY) or "").strip())
        if new_only and has_report_id:
            continue
        # 未 force 时：reportId + AI 都齐才跳过；缺 AI 的会进入补齐轮询
        if not force and has_report_id and has_ai:
            continue
        selected.append((index, row))
        if limit > 0 and len(selected) >= limit:
            break
    return selected


def run_from_csv(args: argparse.Namespace) -> int:
    _apply_accounts_csv(args)
    csv_path = Path(args.from_csv)
    if not csv_path.is_absolute():
        csv_path = Path(__file__).with_name(csv_path.name) if not csv_path.exists() else csv_path
    if not csv_path.exists():
        raise FileNotFoundError(f"找不到 CSV：{csv_path}")

    type_filter = {
        part.strip()
        for part in (args.types or "").split(",")
        if part.strip()
    }
    invalid = type_filter - set(CATEGORY_BY_NAME)
    if invalid:
        raise ValueError(f"--types 仅支持 日报/周报/月报，收到：{sorted(invalid)}")

    fieldnames, rows = load_report_csv(csv_path)
    targets = filter_csv_rows(
        rows,
        type_filter=type_filter,
        limit=args.limit,
        force=args.force,
        new_only=bool(getattr(args, "new_only", False)),
    )
    if not targets:
        print(f"[SKIP] 没有待发送行（csv={csv_path} types={sorted(type_filter) or '全部'}）")
        return 0

    print(f"[CSV] {csv_path.name} 待处理 {len(targets)} 行")
    for index, row in targets:
        type_name = row["周报类型"].strip()
        category = CATEGORY_BY_NAME[type_name]
        start_day, end_day, label = parse_csv_period(row["日期"], category)
        preview = (row.get("周报内容") or "").replace("\n", " ")[:60]
        print(
            f"  #{index + 1} {type_name} {row['日期']} "
            f"{row['发送人']} -> {row['接收人']} "
            f"period={start_day}~{end_day} content={preview}..."
        )

    if not args.execute:
        print("确认后使用 --execute（可加 --poll）实际发送。")
        return 0

    auth_cache: dict[tuple[str, str], dict[str, str]] = {}
    success = 0
    for seq, (index, row) in enumerate(targets, start=1):
        type_name = row["周报类型"].strip()
        category = CATEGORY_BY_NAME[type_name]
        sender_name = row["发送人"].strip()
        receiver_name = row["接收人"].strip()
        try:
            start_day, end_day, label = parse_csv_period(row["日期"], category)
            existing_report_id = (row.get(CSV_RESULT_REPORT_ID) or "").strip()
            existing_ai = (row.get(CSV_RESULT_AI_SUMMARY) or "").strip()

            sender = login_as_user(
                user_name=sender_name,
                company_name=args.company_name,
                login_base_url=args.login_base_url,
                cache=auth_cache,
            )
            receiver = login_as_user(
                user_name=receiver_name,
                company_name=args.company_name,
                login_base_url=args.login_base_url,
                cache=auth_cache,
            )

            # 已有 reportId、只缺 AI：直接补齐轮询，不再重复发报
            if existing_report_id and not existing_ai and not args.force:
                report_id = existing_report_id
                print(
                    f"[1/3] 已有 reportId={report_id}，跳过发报，直接补齐 AI 摘要 "
                    f"({type_name} {label})",
                    flush=True,
                )
            else:
                sender_headers = build_headers(
                    token=sender["token"],
                    tenant_id=sender["tenant_id"],
                    company_id=sender["company_id"],
                    user_id=sender["user_id"],
                )

                # [1/3] 发报：create_report_data 两步创建（草稿拿 ID → 正式提交）
                draft_id, _ = allocate_report_id(sender_headers, category)
                payload = {
                    "reportId": draft_id,
                    "category": category,
                    "content": build_content_from_text(row.get("周报内容") or ""),
                    "reportTo": [
                        {
                            "userId": receiver["user_id"],
                            "name": receiver_name,
                        }
                    ],
                    "ccTo": [],
                    "startDate": start_of_day_ms(start_day),
                    "status": 1,
                }
                print(
                    f"[1/3] 发送{type_name}: {sender_name}({sender['user_id']}) -> "
                    f"{receiver_name}({receiver['user_id']}) period={label}",
                    flush=True,
                )
                response = post_report(sender_headers, payload)
                if not response_ok(response):
                    raise RuntimeError(f"提交失败：{response}")

                # reportId 以发报接口返回为准，缺省再回退草稿 ID
                report_id = extract_report_id(response) or draft_id
                if not report_id:
                    raise RuntimeError(f"发报成功但未返回 reportId：{response}")
                rows[index][CSV_RESULT_REPORT_ID] = report_id
                # 发报成功立刻回写 reportId，避免轮询中断丢结果
                write_report_csv(csv_path, fieldnames, rows)
                print(f"    reportId = {report_id}（已写入 {csv_path.name}）", flush=True)

            ai_text = ""
            # CSV 执行默认轮询 AI 摘要；显式 --no-poll 可关闭
            do_poll = not bool(getattr(args, "no_poll", False))
            if do_poll:
                print(
                    f"[2/3] 用接收方轮询详情 (每 {args.poll_interval:g}s, "
                    f"最多 {args.max_polls} 次) reportId={report_id}",
                    flush=True,
                )
                ai_text, _, ai_summary = poll_ai_summary(
                    receiver_auth=receiver,
                    report_id=report_id,
                    poll_interval=args.poll_interval,
                    max_polls=args.max_polls,
                )
                rows[index][CSV_RESULT_AI_SUMMARY] = ai_text
                write_report_csv(csv_path, fieldnames, rows)
                print("[3/3] 摘要已生成 ✓（已写入 AI总结内容）", flush=True)
                print_ai_summary_like_original(ai_summary)
            else:
                print("[2/3] 已跳过 AI 摘要轮询（--no-poll）", flush=True)

            success += 1
            print(
                f"[OK] {seq}/{len(targets)} {type_name} "
                f"period={label} reportId={report_id} "
                f"ai_len={len(ai_text)}",
                flush=True,
            )
        except Exception as exc:
            print(
                f"[FAILED] {seq}/{len(targets)} row={index + 1} {type_name} error={exc}",
                flush=True,
            )
            write_report_csv(csv_path, fieldnames, rows)
            if not args.continue_on_error:
                raise

    print(f"[DONE] success={success}/{len(targets)} wrote {csv_path}", flush=True)
    return 0 if success == len(targets) else 1


def _apply_accounts_source(args: argparse.Namespace) -> None:
    import excel_tool

    raw_xlsx = (getattr(args, "accounts_xlsx", "") or getattr(args, "accounts_csv", "") or "").strip()
    if raw_xlsx:
        excel_tool.DATA_XLSX_PATH = os.path.abspath(raw_xlsx)
        excel_tool.DATA_CSV_PATH = excel_tool.DATA_XLSX_PATH
        excel_tool._ACCOUNTS_CACHE.clear()
    sheet = (getattr(args, "accounts_sheet", "") or "").strip()
    if sheet:
        os.environ["REPORT_AGENT_DATA_SHEET"] = sheet


def _apply_accounts_csv(args: argparse.Namespace) -> None:
    _apply_accounts_source(args)


def _prescan_env_from_argv() -> str:
    import sys

    env = (os.environ.get("REPORT_AGENT_ENV") or "test").strip().lower()
    argv = sys.argv[1:]
    for i, arg in enumerate(argv):
        if arg == "--env" and i + 1 < len(argv):
            return argv[i + 1].strip().lower()
        if arg.startswith("--env="):
            return arg.split("=", 1)[1].strip().lower()
    return env


def _dedupe_report_to(items: list[dict[str, str]]) -> list[dict[str, str]]:
    seen: set[str] = set()
    out: list[dict[str, str]] = []
    for item in items:
        uid = str(item.get("userId") or "").strip()
        if not uid or uid in seen:
            continue
        seen.add(uid)
        out.append({"userId": uid, "name": str(item.get("name") or "").strip() or uid})
    return out


def run_resend_existing(args: argparse.Namespace) -> int:
    """用已有 reportId 重新编辑提交，把额外接收人加入 reportTo（保留原接收人）。"""
    _apply_accounts_csv(args)
    csv_path = Path(args.from_csv)
    if not csv_path.is_absolute():
        csv_path = Path(__file__).with_name(csv_path.name) if not csv_path.exists() else csv_path
    if not csv_path.exists():
        raise FileNotFoundError(f"找不到 CSV：{csv_path}")

    extra_user_id = (args.extra_recipient_user_id or "").strip()
    extra_name = (args.extra_recipient_name or "").strip() or extra_user_id
    if not extra_user_id:
        raise ValueError("请指定 --extra-recipient-user-id（例如 ouvobtsophngdv）")

    type_filter = {
        part.strip()
        for part in (args.types or "").split(",")
        if part.strip()
    }
    sender_prefix = (args.sender_prefix or "智本_").strip()

    fieldnames, rows = load_report_csv(csv_path)
    targets: list[tuple[int, dict[str, str]]] = []
    for index, row in enumerate(rows):
        report_id = (row.get(CSV_RESULT_REPORT_ID) or "").strip()
        if not report_id:
            continue
        type_name = (row.get("周报类型") or "").strip()
        if type_filter and type_name not in type_filter:
            continue
        sender_name = (row.get("发送人") or "").strip()
        if sender_prefix and not sender_name.startswith(sender_prefix):
            continue
        targets.append((index, row))

    if args.limit > 0:
        targets = targets[: args.limit]

    print(
        f"[RESEND] csv={csv_path.name} rows={len(targets)} "
        f"extra={extra_name}({extra_user_id}) company={args.company_name}",
        flush=True,
    )
    for index, row in targets[:5]:
        print(
            f"  preview #{index + 1} {row['周报类型']} {row['日期']} "
            f"{row['发送人']} -> {row['接收人']} reportId={row.get(CSV_RESULT_REPORT_ID)}",
            flush=True,
        )
    if len(targets) > 5:
        print(f"  ... 另有 {len(targets) - 5} 条", flush=True)

    if not args.execute:
        print("确认后加 --execute 实际重发。")
        return 0

    auth_cache: dict[tuple[str, str], dict[str, str]] = {}
    success = 0
    for seq, (index, row) in enumerate(targets, start=1):
        type_name = row["周报类型"].strip()
        category = CATEGORY_BY_NAME[type_name]
        sender_name = row["发送人"].strip()
        receiver_name = row["接收人"].strip()
        report_id = (row.get(CSV_RESULT_REPORT_ID) or "").strip()
        try:
            start_day, _end_day, label = parse_csv_period(row["日期"], category)
            sender = login_as_user(
                user_name=sender_name,
                company_name=args.company_name,
                login_base_url=args.login_base_url,
                cache=auth_cache,
            )
            receiver = login_as_user(
                user_name=receiver_name,
                company_name=args.company_name,
                login_base_url=args.login_base_url,
                cache=auth_cache,
            )
            headers = build_headers(
                token=sender["token"],
                tenant_id=sender["tenant_id"],
                company_id=sender["company_id"],
                user_id=sender["user_id"],
            )
            payload = {
                "reportId": report_id,
                "category": category,
                "content": build_content_from_text(row.get("周报内容") or ""),
                "reportTo": _dedupe_report_to(
                    [
                        {"userId": receiver["user_id"], "name": receiver_name},
                        {"userId": extra_user_id, "name": extra_name},
                    ]
                ),
                "ccTo": [],
                "startDate": start_of_day_ms(start_day),
                "status": 1,
            }
            response = post_report(headers, payload)
            if not response_ok(response):
                raise RuntimeError(f"重发失败：{response}")
            success += 1
            print(
                f"[OK] {seq}/{len(targets)} reportId={report_id} "
                f"{sender_name}->{receiver_name}+{extra_name} period={label}",
                flush=True,
            )
            if args.sleep:
                time.sleep(args.sleep)
        except Exception as exc:
            print(
                f"[FAILED] {seq}/{len(targets)} row={index + 1} "
                f"reportId={report_id} {sender_name} error={exc}",
                flush=True,
            )
            if not args.continue_on_error:
                return 1

    print(f"[RESEND DONE] success={success}/{len(targets)}", flush=True)
    return 0 if success == len(targets) else 1


def run_seed_mode(args: argparse.Namespace) -> int:
    categories = {int(part.strip()) for part in args.categories.split(",") if part.strip()}
    if not categories.issubset({CATEGORY_DAILY, CATEGORY_WEEKLY, CATEGORY_MONTHLY}):
        raise ValueError("--categories 仅支持 1,2,3")
    seeds = select_seeds(categories, args.count)
    today = datetime.now(TOKYO).date()

    plans: list[dict[str, Any]] = []
    for index, item in enumerate(seeds):
        start_day, end_day, label = resolve_period(item, today)
        plans.append(
            {
                "index": index + 1,
                "case_id": item.get("case_id"),
                "category": int(item["category"]),
                "category_name": CATEGORY_NAME[int(item["category"])],
                "startDate": start_of_day_ms(start_day),
                "period_start": start_day.isoformat(),
                "period_end": end_day.isoformat(),
                "period_label": label,
                "summary": item["summary"],
                "seed": item,
            }
        )

    if not args.execute:
        print(f"[DRY-RUN] 将创建 {len(plans)} 条报告（先草稿分配 ID，再提交）。")
        for plan in plans:
            print(
                f"  {plan['index']}. {plan['category_name']}(category={plan['category']}) "
                f"case={plan['case_id']} startDate={plan['startDate']} "
                f"period={plan['period_start']}~{plan['period_end']} "
                f"summary={plan['summary']}"
            )
        print("确认后使用 --execute 实际写入。")
        return 0

    auth = resolve_auth(args)
    headers = build_headers(
        token=auth["token"],
        tenant_id=auth["tenant_id"],
        company_id=auth["company_id"],
        user_id=auth["user_id"],
    )
    created: list[dict[str, Any]] = []
    for plan in plans:
        item = plan["seed"]
        try:
            report_id, draft_response = allocate_report_id(headers, plan["category"])
            print(
                f"[DRAFT] {plan['index']}/{len(plans)} "
                f"{plan['category_name']} reportId={report_id} case={plan['case_id']}"
            )
            if args.draft_only:
                created.append(
                    {
                        "reportId": report_id,
                        "case_id": plan["case_id"],
                        "category": plan["category"],
                        "category_name": plan["category_name"],
                        "period_start": plan["period_start"],
                        "period_end": plan["period_end"],
                        "ok": True,
                        "draft_only": True,
                        "draft_response": draft_response,
                    }
                )
                continue

            payload = {
                "reportId": report_id,
                "category": plan["category"],
                "content": build_content(item, plan["period_label"]),
                "reportTo": [
                    {
                        "userId": args.report_to_user_id or auth["user_id"],
                        "name": args.report_to_name or args.user_name,
                        "avatar": args.report_to_avatar,
                    }
                ],
                "ccTo": [],
                "startDate": plan["startDate"],
                "status": 1,
            }
            response = post_report(headers, payload)
            ok = response_ok(response)
            err_code = response.get("errCode", response.get("errcode", 0))
            created.append(
                {
                    "reportId": report_id,
                    "case_id": plan["case_id"],
                    "category": plan["category"],
                    "category_name": plan["category_name"],
                    "period_start": plan["period_start"],
                    "period_end": plan["period_end"],
                    "ok": ok,
                    "draft_response": draft_response,
                    "response": response,
                }
            )
            status = "CREATED" if ok else "FAILED"
            print(
                f"[{status}] {plan['index']}/{len(plans)} "
                f"{plan['category_name']} reportId={report_id} case={plan['case_id']} "
                f"errCode={err_code}"
            )
            if not ok and not args.continue_on_error:
                raise RuntimeError(
                    f"提交失败 reportId={report_id} errCode={err_code} response={response}"
                )
        except Exception as exc:
            created.append(
                {
                    "case_id": plan["case_id"],
                    "category": plan["category"],
                    "category_name": plan["category_name"],
                    "period_start": plan["period_start"],
                    "period_end": plan["period_end"],
                    "ok": False,
                    "error": str(exc),
                }
            )
            print(f"[FAILED] {plan['index']}/{len(plans)} {plan['category_name']} error={exc}")
            if not args.continue_on_error:
                raise

    success_count = sum(1 for item in created if item.get("ok"))
    output = {
        "login_user": args.user_name,
        "company_name": args.company_name,
        "created_count": len(created),
        "success_count": success_count,
        "created": created,
    }
    Path(args.output).write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"[DONE] success={success_count}/{len(created)} wrote {args.output}")
    return 0 if success_count == len(created) else 1


def parse_args() -> argparse.Namespace:
    refresh_env_config(_prescan_env_from_argv())
    parser = argparse.ArgumentParser(description="批量构造日报/周报/月报测试数据")
    parser.add_argument(
        "--env",
        choices=("test", "pre"),
        default=os.environ.get("REPORT_AGENT_ENV", "test"),
        help="运行环境：test=测试 / pre=预发",
    )
    parser.add_argument("--execute", action="store_true", help="实际调用接口写入数据")
    parser.add_argument("--count", type=int, default=6, help="创建数量（种子模式）")
    parser.add_argument(
        "--categories",
        default="1,2,3",
        help="种子模式类型，逗号分隔：1=日报,2=周报,3=月报",
    )
    parser.add_argument("--company-name", default=DEFAULT_COMPANY_NAME)
    parser.add_argument("--user-name", default=DEFAULT_USER_NAME)
    parser.add_argument("--login-base-url", default=DEFAULT_LOGIN_BASE_URL)
    parser.add_argument(
        "--auth-from-env",
        action="store_true",
        help="从 REPORT_TOKEN/REPORT_TENANT_ID/REPORT_COMPANY_ID/REPORT_USER_ID 读取鉴权",
    )
    parser.add_argument("--report-to-user-id", default=DEFAULT_REPORT_TO_USER_ID)
    parser.add_argument("--report-to-name", default=DEFAULT_REPORT_TO_NAME)
    parser.add_argument("--report-to-avatar", default=DEFAULT_REPORT_TO_AVATAR)
    parser.add_argument(
        "--output",
        default=str(FIXTURES_DIR / "created_reports.json"),
        help="种子模式结果落盘路径（不含 token）",
    )
    parser.add_argument(
        "--continue-on-error",
        action="store_true",
        help="单条失败时继续创建其余报告",
    )
    parser.add_argument(
        "--draft-only",
        action="store_true",
        help="只执行第一步分配 reportId（status=0），不提交正式内容",
    )
    parser.add_argument(
        "--from-csv",
        nargs="?",
        const=str(DEFAULT_REPORT_CSV),
        default="",
        help="从 report_data.csv 读取发送数据；可只写 --from-csv 使用默认文件",
    )
    parser.add_argument(
        "--types",
        default="",
        help="CSV 模式仅处理指定类型，逗号分隔：日报,周报,月报；默认全部",
    )
    parser.add_argument("--limit", type=int, default=0, help="CSV 模式最多处理条数，0=全部")
    parser.add_argument(
        "--force",
        action="store_true",
        help="CSV 模式即使已有 reportId 也重新发送",
    )
    parser.add_argument(
        "--new-only",
        action="store_true",
        help="CSV 模式跳过已有 reportId 的行，只发送尚未发报的新行",
    )
    parser.add_argument(
        "--poll",
        action="store_true",
        help="CSV 模式发送后轮询 AI 摘要（默认已开启；保留此参数兼容旧调用）",
    )
    parser.add_argument(
        "--no-poll",
        action="store_true",
        help="CSV 模式只发报，不轮询 AI 摘要",
    )
    parser.add_argument("--poll-interval", type=float, default=10.0, help="轮询间隔秒")
    parser.add_argument("--max-polls", type=int, default=60, help="最大轮询次数")
    parser.add_argument(
        "--accounts-xlsx",
        default="",
        help="账号 data.xlsx 路径（默认 fixtures/online/data.xlsx）",
    )
    parser.add_argument(
        "--accounts-sheet",
        default="",
        help="账号 sheet（默认随 --env：test_data / pre_data）",
    )
    parser.add_argument(
        "--accounts-csv",
        default="",
        help=argparse.SUPPRESS,
    )
    parser.add_argument(
        "--resend-existing",
        action="store_true",
        help="用已有 reportId 重新编辑提交（需配合 --from-csv）",
    )
    parser.add_argument(
        "--extra-recipient-user-id",
        default="ouvobtsophngdv",
        help="额外汇报对象 userId，例如可大力",
    )
    parser.add_argument(
        "--extra-recipient-name",
        default="可大力",
        help="额外汇报对象显示名",
    )
    parser.add_argument(
        "--sender-prefix",
        default="智本_",
        help="仅重发发送人前缀匹配的汇报，默认智本_",
    )
    parser.add_argument(
        "--sleep",
        type=float,
        default=0.15,
        help="重发时每条之间的间隔秒，默认 0.15",
    )
    parser.add_argument(
        "--split-csv",
        nargs="?",
        const=str(DEFAULT_REPORT_CSV),
        default="",
        help="按发送人拆分：report_data/{发送人}.csv + report_data.xlsx（每发送人一个 sheet）",
    )
    parser.add_argument(
        "--export-xlsx",
        nargs="?",
        const=str(DEFAULT_REPORT_CSV),
        default="",
        help="仅从 report_data.csv 生成 report_data.xlsx（sheet 名=发送人）",
    )
    parser.add_argument(
        "--merge-csv",
        nargs="?",
        const=str(DEFAULT_REPORT_CSV_DIR),
        default="",
        help="把 report_data/*.csv 合并回 report_data.csv",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    refresh_env_config(args.env)
    _apply_accounts_source(args)
    if args.export_xlsx:
        export_report_data_xlsx(args.export_xlsx)
        return 0
    if args.split_csv:
        split_report_csv_by_sender(args.split_csv)
        return 0
    if args.merge_csv:
        merge_report_csv_dir(args.merge_csv)
        return 0
    if args.from_csv:
        if args.resend_existing:
            return run_resend_existing(args)
        return run_from_csv(args)
    return run_seed_mode(args)


if __name__ == "__main__":
    raise SystemExit(main())
